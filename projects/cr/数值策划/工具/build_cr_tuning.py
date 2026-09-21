"""TASK-0036历史候选重现：A/B等级拉伸已被User撤销，不可用于当前dev提交。

当前等级使用build_cr_level_dev.py（已有1–300逐级对标、后续拟合）；VIP暂存。
以下固定r7013旧候选仅供历史追溯，源表只读，完整值仅输出受控目录。

python build_cr_tuning.py extract|verify --baseline <原受控根> --output <本轮受控目录>
XLSX由render_cr_tuning.mjs生成；未选定的等级映射和未确认真相源的价格表不写值。
"""
from __future__ import annotations
import argparse
import csv
import json
import math
import re
from pathlib import Path
import shutil
from zipfile import ZipFile
import xml.etree.ElementTree as ET
import openpyxl
from producer_workbook_sources import read_source
from build_cashfrenzy_curves import sanitize

FILE = 'CR_调优候选_vs_CF_POP_数值曲线.xlsx'
SHEETS = ['候选总览','VIP_消费门槛','VIP_商城金币','等级_升级消耗','等级_Bet曲线','等级_消耗返还','档位_膨胀','等级_膨胀']
TARGET = [0,94,375,2000,6875,27500,75000,312500,937500,2500000]
POP = [0,75,300,1600,5500,22000,60000,250000,750000,2000000]


def dump(path: Path, value: object) -> None:
    path.write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf-8')


def prepare(base: Path, output: Path) -> dict:
    src = base/'producer-dashboard-20260920/source/r7013'
    receipt = json.loads((base/'producer-dashboard-20260920/export-receipt.json').read_text(encoding='utf-8-sig'))
    assert receipt['revision'] == 7013 and receipt['environment'] == 'trunk'
    specs = {r['table']:r['file'] for r in receipt['files']}
    def table(name: str) -> list[dict]:
        return read_source(src/specs[name])['records']
    vip, levels, ps, prices, bets, unlocks = [table(n) for n in ['VipCfg','LevelCfg','PriceSetting','PriceCheatSheet','SlotsCasinoBetList','SlotsCasinoBetUnlock']]
    # Reuse source snapshots from the last verified merged workbook; do not rerun accepted models.
    merged = openpyxl.load_workbook(base/'outputs/task0036-cr-cf-fixed-vip-20260921/CR_vs_CashFrenzy_数值曲线对照.xlsx',read_only=True,data_only=True)
    cr_cost = list(merged['等级_升级消耗'].iter_rows(min_row=32,max_row=5030,max_col=7,values_only=True))
    cr_bet = list(merged['等级_Bet曲线'].iter_rows(min_row=32,max_row=5031,max_col=11,values_only=True))
    merged.close()
    assert len(levels)==5000 and [r['level'] for r in levels]==list(range(1,5001))
    old = openpyxl.load_workbook(base/'CashRoyal数值.xlsx',read_only=True,data_only=True)
    tier = [list(r)[:2] for r in list(old['档位膨胀'].values)[1:] if isinstance(r[0],(int,float))]
    cf = [list(r) for r in old['cashFrenzy等级'].iter_rows(min_row=2,max_row=301,max_col=9,values_only=True)]
    old.close()
    # Bound the mapping search to the five formal growth/denomination logic sheets.
    old = openpyxl.load_workbook(base/'CashRoyal数值.xlsx',read_only=True,data_only=False)
    mapping_hits=[]
    for name in ['等级+buff','旧cr等级','等级体验表','档位膨胀','基础金币']:
        for row in old[name]:
            for c in row:
                value=c.value.text if hasattr(c.value,'text') else c.value
                if isinstance(value,str) and value.startswith('=') and 'cashfrenzy' in value.lower():
                    mapping_hits.append([name,c.coordinate,value])
    old.close()
    assert not mapping_hits, '历史映射证据变化，先复核，不自动选候选'
    # POP Tier6–10 log slope constrained to the Tier10 anchor. Only the tail is fitted.
    slope=sum((k-10)*math.log(POP[k-1]/POP[9]) for k in range(6,11))/sum((k-10)**2 for k in range(6,11))
    target = TARGET + [math.floor(TARGET[-1]*math.exp(slope*(v-10))+.5) for v in range(11,16)]
    assert all(a<b for a,b in zip(target,target[1:])) and target[-1] < 2**31
    vip_rows=[[r['vipLevel'],r['needExp'],target[i],POP[i] if i<10 else 'N/A',r['_excel_row']] for i,r in enumerate(vip)]
    # Match the historical denomination shape to every current shop-price row.
    shops=sorted((r for r in ps if (r['currencyType'],r['vipType'])==(1,1)),key=lambda r:r['money'])
    assert [r['money'] for r in shops]==[r[0] for r in tier] and len(tier)==30
    base_coin=shops[0]['vip0'];factors=dict(tier)
    for r in shops:
        assert math.isclose(r['vip0']/((r['money']+1)/100)/base_coin,factors[r['money']])
    assert factors[tier[-1][0]]>factors[tier[0][0]]
    tier_rows=[]
    for i,r in enumerate(shops):
        f=factors[r['money']];normal=1+4*(f-tier[0][1])/(tier[-1][1]-tier[0][1])
        affected=[p for p in ps if p['money']==r['money'] and p['currencyType']==1]
        tier_rows.append([i+1,r['money'],f,r['vip0'],normal,r['vip0']*normal/f,len(affected),r['_excel_row']])
    assert tier_rows[0][4]==1 and tier_rows[-1][4]==5
    assert all(a[4]<=b[4] for a,b in zip(tier_rows,tier_rows[1:]))
    proposal=[]
    for r in ps:
        if r['currencyType']!=1:continue
        t=next(x for x in tier_rows if x[1]==r['money']);scale=t[4]/t[2]
        for v in range(16):
            val=r[f'vip{v}'];new=val*scale
            assert math.isclose(new,round(new),abs_tol=1e-5)
            assert math.isclose(round(new)/round(r['vip0']*scale),val/r['vip0'])
            if round(new)!=val:proposal.append(['PriceSetting.xlsx','Sheet1',f'{openpyxl.utils.get_column_letter(v+4)}{r["_excel_row"]}',r['money'],f'vip{v}',val,round(new),'仅方案，未写配置'])
    # Level inflation is a value multiplier, independent of the level-up cost target.
    basis=[r for r in prices if (r['money'],r['priceType'],r['vipType'])==(100,17,1)]
    by_bet={r['levelId']:r for r in bets};level_rows=[]
    for i,r in enumerate(levels):
        lv=r['level'];unlock=max((u for u in unlocks if u['roleLevel']<=lv),key=lambda u:u['betlevel']) if 'roleLevel' in unlocks[0] else None
        if unlock is None:
            from producer_dashboard_models import max_unlocked
            unlock=max_unlocked(unlocks,lv)
        exp=by_bet[unlock['betlevel']]['levelExp']
        price=max((p for p in basis if p['level']<=lv),key=lambda p:p['level'])
        old_cost=cr_cost[i][1] if i<4999 else 'N/A';reward=cr_cost[i][3] if i<4999 else 'N/A'
        cb=cr_bet[i]
        if i<300:assert math.isclose(price['vip_16_0']/base_coin,cf[i][8])
        level_rows.append([lv,r['levelUpType'] if r['levelUpType'] is not None else 0,r['levelUpExp'],exp,cb[9],old_cost,reward,cb[1],cb[2],price['vip_16_0']/base_coin,cf[i][8] if i<300 else 'N/A'])
    # PriceSetting has no level key; current price service resolves through PriceCheatSheet.
    ev=output/'evidence'
    legacy=(ev/'PriceSettingSvc.cs').read_text(encoding='utf-8-sig')
    current=(ev/'PriceCheatSheetSvc.cs').read_text(encoding='utf-8-sig')
    assert '此类已废弃' in legacy and 'PriceCheatSheetDescMgr.Instance.GetValue' in current
    code=base/'producer-dashboard-20260920/business-code-r7013'
    zero=(code/'VipLevelUtils.cs').read_text(encoding='utf-8-sig')
    login=(ev/'LoginHandler.cs').read_text(encoding='utf-8-sig')
    assert 'if (vipExp == 0)' in zero and 'desc.Value.needExp > player.RoleBase.VipExp' in zero
    assert 'VipLevelUtils.CorrectVipExp(player)' in login
    def resolve(exp: int) -> int:
        return next((i for i,t in enumerate(target) if t>exp),15)
    boundary=[{'exp':x,'level':resolve(x)} for x in [0,1,93,94,374,375,2499999,2500000,target[-1]]]
    assert resolve(0)==1 and resolve(94)==2 and resolve(2500000)==10
    dump(output/'vip-zero-check.json',{'status':'Not Passed / candidate only','static_bounded_loop':True,'login_zero_exp_level':1,'AddVipExp_zero_returns_without_correction':True,'AddVipLevel_0_to_1_delta':0,'boundaries':boundary,'client_and_server_runtime':'Not run; no initialized runtime evidence'})
    for name in ['LevelCfg','PriceSetting']:
        dest=output/'withheld'/f'{name}.xlsx';dest.parent.mkdir(exist_ok=True)
        shutil.copyfile(src/specs[name],dest)
    (output/'candidate').mkdir(exist_ok=True)
    (output/'diff').mkdir(exist_ok=True)
    with (output/'diff/PriceSetting_proposal_NOT_APPLIED.csv').open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f);w.writerow(['file','sheet','cell','money','field','before','proposal','status']);w.writerows(proposal)
    data={'revision':7013,'source_vip':str(src/specs['VipCfg']),'source_dir':str(src),'specs':specs,'vip':vip_rows,'pop':POP,'tail_slope':slope,'tail_growth':math.exp(slope),'levels':level_rows,'cf':[[r[0],r[6],r[5],r[8]] for r in cf],'tiers':tier_rows,'shop_vip':[[v,shops[0][f'vip{v}']/base_coin] for v in range(16)],'price_proposal_cells':len(proposal),'mapping':'No unique formal mapping; A full stretch / B first100 retained','gates':{'VIP':'拟合候选；零门槛安全未通过','LevelCfg':'无唯一映射；两套体验候选，不写表','PriceSetting':'运行查询已转向PriceCheatSheet；未证实生成关系，不写表','level_inflation':'1–300一致 No Change；301–5000缺CF证据'}}
    dump(output/'tuning-inputs.json',data)
    return {'status':'prepared','vip_rows':15,'level_rows':5000,'business_tiers':30,'price_proposal_cells':len(proposal)}


def verify(base: Path, output: Path) -> dict:
    d=json.loads((output/'tuning-inputs.json').read_text(encoding='utf-8'))
    changes=[]
    for name,folder in [('VipCfg','candidate'),('LevelCfg','withheld'),('PriceSetting','withheld')]:
        old=openpyxl.load_workbook(Path(d['source_dir'])/d['specs'][name],read_only=True,data_only=False)
        new=openpyxl.load_workbook(output/folder/(name+'.xlsx'),read_only=True,data_only=False)
        assert old.sheetnames==new.sheetnames
        for a,b in zip(old,new):
            ar=list(a.values);br=list(b.values)
            assert len(ar)==len(br)
            for r,(x,y) in enumerate(zip(ar,br),1):
                assert len(x)==len(y)
                for c,(before,after) in enumerate(zip(x,y),1):
                    if before==after:continue
                    assert name=='VipCfg' and a.title=='Sheet1' and c==2 and 5<=r<=19
                    assert after==d['vip'][r-5][2]
                    changes.append([name+'.xlsx',a.title,f'B{r}',r-4,'needExp',before,after,after-before,'候选；零门槛Gate未通过'])
        old.close();new.close()
    assert len(changes)==15
    with (output/'diff/config_actual_diff.csv').open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f);w.writerow(['file','sheet','cell','id','field','before','after','delta','status']);w.writerows(changes)
    f=openpyxl.load_workbook(output/FILE,data_only=False)
    v=openpyxl.load_workbook(output/FILE,data_only=True)
    assert [s.title for s in f if s.sheet_state=='visible']==SHEETS
    def same(a: object,b: object) -> None:
        assert (isinstance(a,(int,float)) and math.isclose(a,b,rel_tol=1e-10,abs_tol=1e-7)) if isinstance(b,(int,float)) else a==b,(a,b)
    for i,row in enumerate(d['vip'],32):
        for col,want in zip('BCDE',[row[1]/100,row[2]/100,POP[row[0]-1]/80 if row[0]<=10 else 'N/A',row[2]]):same(v['VIP_消费门槛'][f'{col}{i}'].value,want)
    for i,row in enumerate(d['levels'],2):
        lv,mode,oldexp,exp,bet,cost,reward,rawbet,eqbet,inf,cf_inf=row
        s=v['CALC_等级'];same(s.cell(i,1).value,lv)
        same(v['等级_升级消耗'].cell(i+30,2).value,cost)
        for index_col,target_col,spins_col,need_col,cost_col,return_col,kind in [(2,3,4,5,6,7,'A'),(8,9,10,11,12,13,'B')]:
            if lv==5000:
                same(s.cell(i,cost_col).value,'N/A');continue
            mapped=1+math.floor((lv-1)*299/4998) if kind=='A' else lv if lv<=100 else 101+math.floor((lv-101)*199/4898)
            target=d['cf'][mapped-1][1]
            spins=max(1,math.ceil(target/(bet*.05)-1e-10))
            threshold=spins if mode==1 else math.ceil(spins*exp-1e-10)
            actual_spins=threshold if mode==1 else math.ceil(threshold/exp-1e-10)
            for c,want in [(index_col,mapped),(target_col,target),(spins_col,spins),(need_col,threshold),(cost_col,actual_spins*bet*.05),(return_col,reward/(actual_spins*bet*.05))]:same(s.cell(i,c).value,want)
    assert v['档位_膨胀']['D32'].value==1 and v['档位_膨胀']['D61'].value==5
    assert all(v['档位_膨胀'].cell(r,4).value<=v['档位_膨胀'].cell(r+1,4).value for r in range(32,61))
    errors=[];missing=[]
    for s in f:
        for row in s:
            for c in row:
                cv=v[s.title][c.coordinate]
                if cv.data_type=='e':errors.append([s.title,c.coordinate,cv.value])
                if c.data_type=='f' and cv.value is None:missing.append([s.title,c.coordinate])
    assert not errors and not missing,(errors[:5],missing[:5])
    with ZipFile(output/FILE) as z:
        assert not any(x.startswith('xl/externalLinks/') for x in z.namelist())
        assert not any(re.search(r'(?<![A-Za-z])(?:[A-Za-z]:[\\/]|file:/)',z.read(n).decode()) for n in z.namelist() if n.endswith(('.xml','.rels')))
        core=ET.fromstring(z.read('docProps/core.xml'))
        assert not any(c.text for c in core if c.tag.rsplit('}',1)[-1] in ('creator','lastModifiedBy'))
        ns={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
        charts=sorted(x for x in z.namelist() if re.fullmatch(r'xl/charts/chart\d+\.xml',x))
        assert len(charts)==7
        points=[]
        for name in charts:
            root=ET.fromstring(z.read(name));lines=root.findall('.//c:lineChart',ns)
            assert len(lines)==1 and len(root.findall('.//c:valAx',ns))==1
            points.append([sum(p.text not in (None,'','#N/A') for p in ser.findall('c:val/c:numRef/c:numCache/c:pt/c:v',ns)) for ser in lines[0].findall('c:ser',ns)])
        assert points==[[10,10,10],[16,16],[4999,4999,4999,300],[5000,5000],[4999,4999,4999],[30,30],[300,300]],points
    for name in SHEETS:assert f[name].freeze_panes is None
    result={'task':'TASK-0036','status':'Review / configuration gates open','source_revision':7013,'visible_sheets':len(SHEETS),'line_charts':7,'vip_cells_changed':15,'VIP_first10_exact':True,'VIP11_15_fit':'POP Tier6–10 log-linear with fixed VIP10 anchor','VIP_monotonic':True,'VIP_zero_runtime_gate':'Not Passed; login promotes zero exp to VIP1; client/server runtime not verified','LevelCfg_cells_changed':0,'mapping_candidates':2,'level_rows':5000,'PriceSetting_cells_changed':0,'price_proposal_cells':d['price_proposal_cells'],'denomination_endpoints':[1,5],'shape_preserved':True,'level_inflation_1_300':'No Change, 300 matching levels','level_inflation_301_5000':'No CF source; no extrapolated confirmation','formula_errors':0,'missing_caches':0,'external_links':0,'frozen_sheets':0,'source_written':False,'SVN_committed':False}
    f.close();v.close();dump(output/'validation.json',result)
    return result


def main() -> None:
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('action',choices=['extract','verify','sanitize']);p.add_argument('--baseline',type=Path,required=True);p.add_argument('--output',type=Path,required=True);a=p.parse_args()
    assert not a.output.resolve().is_relative_to(Path(__file__).resolve().parents[4])
    a.output.mkdir(parents=True,exist_ok=True)
    if a.action=='sanitize':
        for folder,name in [('',FILE),('candidate','VipCfg.xlsx'),('withheld','LevelCfg.xlsx'),('withheld','PriceSetting.xlsx')]:sanitize(a.output/folder,name)
        result={'metadata':'sanitized'}
    else:result=prepare(a.baseline,a.output) if a.action=='extract' else verify(a.baseline,a.output)
    print(json.dumps(result,ensure_ascii=False))


if __name__=='__main__':main()
