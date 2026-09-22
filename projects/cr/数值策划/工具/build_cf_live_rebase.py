"""TASK-0036: read-only current CF evidence and r7258 comparison.

prepare --analysis <controlled CF analysis> --base <controlled numerics root> --out <output>
verify --out <output>; render_cf_live_rebase.mjs authors XLSX using Artifact.
No network, SVN operations, source writes, historical CF fallbacks or hashes.
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import re
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET
import openpyxl
from build_cashfrenzy_curves import sanitize

NAMES = ['VIP_消费门槛', 'VIP_膨胀系数', '等级_升级消耗', '等级_Bet曲线', '等级_升级消耗返还']
FILES = ['CR_vs_CF_最新实测_数值曲线.xlsx', 'r7258_vs_CF_最新实测_差异.xlsx']


def csv_rows(root: Path, name: str) -> list[dict]:
    with (root/name).open(encoding='utf-8-sig', newline='') as handle:
        return list(csv.DictReader(handle))


def prepare(analysis: Path, base: Path, out: Path) -> dict:
    levels = csv_rows(analysis, 'cf_levels_measured.csv')
    measured = {int(r['level']): r for r in levels}
    assert list(measured) == list(range(6, 51))
    multis = [[int(r[k]) for k in ('level_from', 'level_to', 'multiplier')]
              for r in csv_rows(analysis, 'cf_road_level_multiplier.csv')]
    assert multis == [[2,9,1],[10,19,2],[20,29,4],[30,39,6],[40,49,8],[50,74,10],[75,99,12],[100,124,14]]
    bonus = {int(r['level']): int(r['bonus_coins']) for r in csv_rows(analysis, 'cf_road_level_bonus.csv')}
    changes = [int(r['level']) for r in csv_rows(analysis, 'cf_road_bet_change_levels.csv')]
    assert changes == [3,4,5,10,15,20,25,30,35,40,50,75,100]
    spins = csv_rows(analysis, 'cf_spins.csv')
    assert len(spins) == 1020
    # Read only these numerical fields; no raw rows or account/log fields copied.
    segments, gain, checked, ratio_conflicts, zero_exp = {}, 0, 0, 0, 0
    for r in spins:
        exp = float(r['xp_gain']) if r['xp_gain'] else 0
        if r['xp_per_bet']:
            ratio = float(r['xp_per_bet'])
            if math.isclose(ratio, 1/3, abs_tol=1e-6):
                checked += 1
            elif ratio == 0:
                zero_exp += 1
            else:
                ratio_conflicts += 1
        gain += exp
        if r['levelup'] == '1':
            segments[int(r['level'])] = gain
            gain = 0
    cf = []
    for level in range(1, 125):
        r = measured.get(level)
        mult = next((m for low, high, m in multis if low <= level <= high), None)
        delta = int(r['xp_delta_from_previous']) if r and r['xp_delta_from_previous'] else None
        if delta is not None:
            assert delta == int(r['xp_at_levelup'])-int(measured[level-1]['xp_at_levelup'])
            assert math.isclose(delta, segments[level], abs_tol=1e-6)
        reward = int(r['level_up_bonus']) if r else None
        if r and level < 50:
            assert reward == bonus[level]
        if level == 50:
            assert reward == 11500000 and bonus[level] == 460000
        cf.append([level, int(r['max_bet_multiplier']) if r else None, mult,
                   delta, reward, int(r['spins_spent']) if r else None,
                   f'cf_levels_measured.csv row {level-4}' if r else '无升级实测',
                   'cf_road_level_multiplier.csv', level in changes])
    observed = [(r[0], r[1]) for i,r in enumerate(cf) if r[1] is not None and (i == 0 or cf[i-1][1] != r[1])]
    assert observed == [(6,300000),(10,450000),(15,1050000),(20,3000000),(25,6000000),(30,7500000),(35,9000000),(40,15000000),(50,30000000)]
    vip_text = (analysis/'CF_VIP_RULES.md').read_text(encoding='utf-8')
    vip_points = [int(x.strip()) for x in re.search(r'VIP_POINTS_THRESHOLD = \{([^}]+)\}', vip_text)[1].split(',')]
    assert vip_points == [0,150,4100,31000,260000,2100000,10000000,50000000]
    packages = [float(v) for v in re.findall(r'^\{\s*([\d.]+),[^\n]+\}, --', vip_text, re.M)]
    assert packages == [1,1.5,2.5,4,7,10,20,40,50]
    prices = [float(x) for x in re.search(r'^\| 价格 \|(.+)\|$', vip_text, re.M)[1].replace('$','').split('|')]
    points = [int(x.replace(',','').strip()) for x in re.search(r'^\| vip_point \|(.+)\|$', vip_text, re.M)[1].split('|')]
    shops = list(zip(prices, points))
    assert len(shops) == 6
    outputs = base/'outputs'
    dev = outputs/'task0036-bet-dev-20260921'
    receipt = json.loads((dev/'svn-result.json').read_text(encoding='utf-8'))
    assert receipt['committed_revision'] == 7258
    old = json.loads((dev/'bet-dev-inputs.json').read_text(encoding='utf-8'))
    tuning = json.loads((outputs/'task0036-tuning-pop-cf-20260921/tuning-inputs.json').read_text(encoding='utf-8'))
    prior = json.loads((outputs/'task0036-cr-cf-fixed-vip-20260921/comparison-inputs.json').read_text(encoding='utf-8'))
    # Reuse actual CR inputs, never old CF or inferred CF values.
    cr = [[r[0],r[1],r[2],c['official_bet'],c['candidate_exp'],r[6],
           prior['cr'][2][i][2] if i < 4999 else None,r[7],r[8],r[9],r[10]]
          for i,(r,c) in enumerate(zip(old['cr'],old['calculations']))]
    assert len(cr) == 5000
    # Bounded check that reused Bet/EXP fields match submitted remote readback.
    b = openpyxl.load_workbook(dev/'remote/SlotsCasinoBetList.xlsx', read_only=True, data_only=True)
    rows = list(b.worksheets[0].iter_rows(values_only=True)); headers = list(rows[0])
    bet_exp = {r[headers.index('bet2')]: r[headers.index('levelExp')] for r in rows[4:] if r[0] is not None}
    b.close()
    assert all(bet_exp[r[3]] == r[4] for r in cr[5:50])
    data = {'cf':cf,'cr':cr,'vip':[[r[0],r[1],r[2]] for r in tuning['vip']],
            'cr_packages':tuning['shop_vip'],'cf_vip':[[i+1,p] for i,p in enumerate(vip_points)],
            'cf_packages':[[i+1,m] for i,m in enumerate(packages)],'shops':shops,
            'drivers':{'base_a':500000,'base_b':150000,'cf_loss':0.15,'cr_loss':old['cr_loss'],
                       'discount':5.5,'cr_vip_points_usd':100},
            'evidence':{'date':'2026-09-22','spins':len(spins),'upgrades':len(levels),
                        'exp_ratio_rows':checked,'exp_ratio_conflicts':ratio_conflicts,'zero_exp_rows':zero_exp,
                        'closed_transitions':44,'segments_gap':0,
                        'cr_revision':7258,'cr_reward_source':'r7013 Accepted result, unchanged by r7258',
                        'level_index':'arrival L7..L50; previous-level Bet and conversion for whole transition'}}
    out.mkdir(parents=True, exist_ok=True)
    (out/'live-inputs.json').write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    return data['evidence']


def verify(out: Path) -> dict:
    d = json.loads((out/'live-inputs.json').read_text(encoding='utf-8'))
    results = []
    for filename in FILES:
        sanitize(out, filename)
        f = openpyxl.load_workbook(out/filename, data_only=False)
        v = openpyxl.load_workbook(out/filename, data_only=True)
        def same(value: object, expected: object) -> None:
            if isinstance(expected, (int, float)):
                assert isinstance(value,(int,float)) and math.isclose(value,expected,rel_tol=1e-10,abs_tol=1e-7),(value,expected)
            else:
                assert value == expected, (value,expected)
        visible = [s.title for s in f if s.sheet_state == 'visible']
        assert visible == (NAMES if filename == FILES[0] else ['r7258差异'])
        for s in f:
            assert s.freeze_panes is None
            for row in s:
                for cell in row:
                    if cell.data_type == 'f':
                        actual = v[s.title][cell.coordinate]
                        assert actual.value is not None and actual.data_type != 'e', (s.title,cell.coordinate,actual.value)
        calc = v['CALC']
        cost_diff_a, cost_diff_b, spin_diff = [], [], []
        for level in range(7,51):
            row = level+3; src = d['cf'][level-1]; before = d['cf'][level-2]; cr = d['cr'][level-2]
            cf_spin = 3*src[3]/before[1]
            cf_gross = cf_spin*before[1]; cf_loss = cf_gross*d['drivers']['cf_loss']
            ca = cf_loss/(d['drivers']['base_a']*before[2]); cb = cf_loss/(d['drivers']['base_b']*before[2])
            ra = src[4]/(d['drivers']['base_a']*before[2]); rb = src[4]/(d['drivers']['base_b']*before[2])
            cr_spin = cr[2] if cr[1] == 1 else cr[2]/cr[4]
            cr_cost = cr_spin*cr[3]/cr[5]*d['drivers']['cr_loss']
            for col,expected in {'F':cf_spin,'G':math.ceil(cf_spin),'H':cf_gross,'I':cf_loss,
                                 'J':ca,'K':cb,'L':ra,'M':rb,'N':ra/ca,'O':rb/cb,
                                 'R':cr_spin,'S':cr_cost}.items():
                same(calc[f'{col}{row}'].value,expected)
            same(calc[f'N{row}'].value,calc[f'O{row}'].value)
            cost_diff_a.append(cr_cost-ca);cost_diff_b.append(cr_cost-cb);spin_diff.append(cr_spin-cf_spin)
        for level in ([1,6,51,75,100,125,300,5000] if filename == FILES[0] else [1,6]):
            for c in 'FGHIJKLMNO':
                same(calc[f'{c}{level+3}'].value,'N/A')
        for level in range(51,125):
            same(v['SRC_CF'].cell(level+3,2).value,'N/A')
        if filename == FILES[0]:
            high = max(p/vv for vv,p in d['shops']); low = min(p/vv for vv,p in d['shops'])
            for vip,pts in d['cf_vip']:
                r = vip+31
                same(v[NAMES[0]].cell(r,4).value,pts/high)
                same(v[NAMES[0]].cell(r,5).value,pts/low)
            assert v[NAMES[0]]['D40'].value == 'N/A'
            same(v[NAMES[1]]['C41'].value,50)
            for name in visible:
                for row in f[name]:
                    for c in row:
                        if c.data_type == 'f':
                            assert not re.search(r'CALC[\x27!]*!?\$?(Y|Z)\$?\d',c.value)
                        assert '5.5' not in str(c.value)
        else:
            for level in range(6,51):
                r=level+26; cr=d['cr'][level-1]; cf=d['cf'][level-1]; t=level+3
                expected=[level,cr[3],cf[1],cr[3]-cf[1],cr[4],cf[1]/3,cr[4]-cf[1]/3,
                          calc[f'R{t}'].value,calc[f'F{t}'].value,
                          calc[f'R{t}'].value-calc[f'F{t}'].value if level>=7 else 'N/A',
                          calc[f'S{t}'].value,calc[f'J{t}'].value,calc[f'K{t}'].value,
                          calc[f'S{t}'].value-calc[f'J{t}'].value if level>=7 else 'N/A',
                          calc[f'S{t}'].value-calc[f'K{t}'].value if level>=7 else 'N/A',
                          cr[5]/d['cr'][0][5],cf[2],cr[5]/d['cr'][0][5]-cf[2],f'{level-1}→{level}']
                for c,want in enumerate(expected,1):
                    same(v['r7258差异'].cell(r,c).value,want)
        ns={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
        with ZipFile(out/filename) as z:
            assert not any(n.startswith('xl/externalLinks/') for n in z.namelist())
            charts=[n for n in z.namelist() if re.fullmatch(r'xl/charts/chart\d+\.xml',n)]
            assert len(charts) == (5 if filename == FILES[0] else 0)
            chart_points=[]
            for name in sorted(charts):
                root=ET.fromstring(z.read(name)); assert len(root.findall('.//c:lineChart',ns)) == 1
                assert len(root.findall('.//c:valAx',ns)) == 1
                chart_points.append([len(s.findall('c:val/c:numRef/c:numCache/c:pt',ns))
                                     for s in root.findall('.//c:ser',ns)])
            assert chart_points == ([[8,8,8,8],[16,9],[44,44,44],[119,45,119,119,119],[44,44]] if filename==FILES[0] else [])
            assert not any(re.search(r'(?<![A-Za-z])(?:[A-Za-z]:[\\/]|file:/)',z.read(n).decode()) for n in z.namelist() if n.endswith(('.xml','.rels')))
        results.append({'file':filename,'visible':visible,'charts':len(charts),'chart_points':chart_points,
                        'formula_errors':0,'external_links':0,'freeze_panes':0})
        f.close();v.close()
    same_bet=sum(a[3]==b[1] for a,b in zip(d['cr'][5:50],d['cf'][5:50]))
    same_exp=sum(math.isclose(a[4],b[1]/3) for a,b in zip(d['cr'][5:50],d['cf'][5:50]))
    same_mult=sum(math.isclose(a[5]/d['cr'][0][5],b[2]) for a,b in zip(d['cr'][5:50],d['cf'][5:50]))
    result={'status':'Review','files':results,'closed_cf_arrivals':'L7-L50','unknown_cf_cost':'L6 and 51+',
            'scenario_return_rates_equal':True,'historical_cf_rows_used':0,'discount_in_visible_or_charts':False,
            'diff_rows':45,'equal_bet_levels':same_bet,'equal_exp_levels':same_exp,'equal_inflation_levels':same_mult,
            'equal_standard_spin_transitions':sum(math.isclose(x,0,abs_tol=1e-9) for x in spin_diff),
            'equal_cost_transitions_a':sum(math.isclose(x,0,abs_tol=1e-9) for x in cost_diff_a),
            'equal_cost_transitions_b':sum(math.isclose(x,0,abs_tol=1e-9) for x in cost_diff_b),
            'cr_cost_lower_transitions_a':sum(x<0 for x in cost_diff_a),
            'cr_cost_lower_transitions_b':sum(x<0 for x in cost_diff_b),
            'exp_rule_status':'Task rule used; raw aggregate contradicts universal empirical claim',
            'exp_ratio_conflicts':d['evidence']['exp_ratio_conflicts'],
            'svn_writes':0,'subagents':'none'}
    (out/'validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    return result


if __name__ == '__main__':
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('action',choices=['prepare','verify']);p.add_argument('--analysis',type=Path)
    p.add_argument('--base',type=Path);p.add_argument('--out',required=True,type=Path)
    args=p.parse_args()
    print(json.dumps(prepare(args.analysis,args.base,args.out) if args.action=='prepare' else verify(args.out),ensure_ascii=False))
