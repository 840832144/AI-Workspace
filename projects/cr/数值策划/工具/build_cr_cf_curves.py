"""TASK-0036 合并曲线：读取已交付缓存，不重算CR/CF既有模型。

extract/verify --cr <CR曲线v2.xlsx> --cf <CF曲线.xlsx> --output <受控目录>
--sgd-usd 仅接受User明确指定值；省略则空白，不查Web、不默认1。
作者为Artifact JS；本脚本的openpyxl只读。sanitize只清理输出元数据。
"""
from __future__ import annotations
import argparse
import json
import math
import re
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET
import openpyxl
from build_cashfrenzy_curves import NAMES, sanitize

FILE = 'CR_vs_CashFrenzy_数值曲线对照.xlsx'
FIRST = 32


def read_cache(path: Path) -> list[list]:
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    result = []
    for i, name in enumerate(NAMES):
        width = [3, 4, 4, 6, 4][i]
        rows = [list(r) for r in book[name].iter_rows(min_row=6, max_col=width, values_only=True)
                if isinstance(r[0], (int, float))]
        result.append(rows)
    book.close()
    return result


def extract(cr: Path, cf: Path, rate: float | None) -> dict:
    """只复制既有结果；新计算仅限本轮授权的CF VIP边界/汇率。"""
    repo = Path(__file__).resolve().parents[4]
    spec = repo/'projects/cr/REPORTS/CR-20260922-PRODUCER-EXPERIENCE/CR_CF_CURVE_COMPARISON.md'
    text = spec.read_text(encoding='utf-8')
    points = [[int(a), int(b.replace(',', ''))] for a, b in re.findall(r'^\| ([0-7]) \| ([\d,]+) \|$', text, re.M)]
    high = float(re.search(r'最高兑换效率：([\d.]+)', text)[1])
    low = float(re.search(r'最低兑换效率：([\d.]+)', text)[1])
    shops = [[float(a), int(b.replace(',', ''))] for a, b in re.findall(r'^\| S\$([\d.]+) \| ([\d,]+) \|', text, re.M)]
    assert len(points) == 8 and len(shops) == 9
    assert round(min(p/v for v, p in shops), 4) == low
    assert round(max(p/v for v, p in shops), 4) == high
    cr_rows, cf_rows = read_cache(cr), read_cache(cf)
    assert [len(r) for r in cr_rows] == [15, 15, 4999, 5000, 4999]
    assert [len(r) for r in cf_rows] == [8, 8, 300, 300, 300]
    for rows in cr_rows + cf_rows[2:]:
        assert [r[0] for r in rows] == list(range(1, len(rows)+1))
    assert sum(r[1] == 'N/A' for r in cf_rows[3]) == 56
    assert rate is None or rate > 0
    return {'cr_file': cr.name, 'cf_file': cf.name, 'cr': cr_rows, 'cf': cf_rows[2:],
            'vip': [[v, p, {2:'黄金',3:'铂金',4:'钛金',5:'尊徽',6:'百夫长',7:'王者风范'}.get(v, '')]
                    for v, p in points], 'shops': shops, 'high': high, 'low': low,
            'fx': rate, 'evidence': 'User 2026-09-21确认及CR_CF_CURVE_COMPARISON.md；本轮未独立读取原截图，不含账号当前点数'}


def verify(data: dict, output: Path) -> dict:
    f = openpyxl.load_workbook(output/FILE, data_only=False)
    v = openpyxl.load_workbook(output/FILE, data_only=True)
    assert [s.title for s in f if s.sheet_state == 'visible'] == NAMES
    assert f['SRC_CR'].sheet_state == f['SRC_CF'].sheet_state == 'hidden'
    fx = v[NAMES[0]]['B4'].value
    assert fx == data['fx']
    def same(got: object, expected: object) -> None:
        if isinstance(expected, (float, int)):
            assert isinstance(got, (float, int)) and math.isclose(got, expected, rel_tol=1e-10, abs_tol=1e-8), (got, expected)
        else:
            assert got == expected, (got, expected)
    for i in range(5):
        for j, row in enumerate(data['cr'][i]):
            r = FIRST+j
            cf = data['cf'][i-2][j] if i >= 2 and j < 300 else None
            if i == 0:
                vip = data['vip'][j+1] if j < 7 else None
                expected = [row[0], row[1], vip[2] or '名称未提供' if vip else 'N/A', vip[1] if vip else 'N/A',
                            vip[1]/data['high'] if vip else 'N/A', vip[1]/data['low'] if vip else 'N/A',
                            vip[1]/data['high']*fx if vip and fx else ('待汇率' if vip else 'N/A'),
                            vip[1]/data['low']*fx if vip and fx else ('待汇率' if vip else 'N/A'),
                            data['high'] if vip else 'N/A', data['low'] if vip else 'N/A', row[2]]
            elif i == 1:
                vip = data['vip'][j+1] if j < 7 else None
                expected = [row[0], row[1], vip[1]/data['vip'][1][1] if vip else 'N/A', row[2], vip[1] if vip else 'N/A']
            elif i == 2:
                expected = [row[0], row[1], cf[1] if cf else 'N/A', row[2], cf[2] if cf else 'N/A', row[3], cf[3] if cf else 'N/A']
            elif i == 3:
                expected = [row[0], row[1], row[3], cf[1] if cf else 'N/A', cf[3] if cf else 'N/A',
                            row[2], row[4], cf[2] if cf else 'N/A', cf[4] if cf else 'N/A',row[5],cf[5] if cf else 'N/A']
            else:
                expected = [row[0],row[1],cf[1] if cf else 'N/A',row[2],cf[2] if cf else 'N/A',row[3],cf[3] if cf else 'N/A']
            for c, want in enumerate(expected, 1):
                same(v[NAMES[i]].cell(r, c).value, want)
    errors, blanks, expected_na = [], [], []
    for s in f:
        for row in s:
            for c in row:
                if c.data_type != 'f':
                    continue
                actual = v[s.title][c.coordinate]
                if actual.value is None:
                    blanks.append((s.title, c.coordinate))
                if actual.data_type == 'e':
                    if s.title == NAMES[0] and c.column in (25, 26) and 32 <= c.row <= 38 and fx is None and actual.value == '#N/A':
                        expected_na.append(c.coordinate)
                    else:
                        errors.append((s.title,c.coordinate))
    assert not errors and not blanks, (errors[:5], blanks[:5])
    assert len(expected_na) == (14 if fx is None else 0)
    ns = {'c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
    points = []
    with ZipFile(output/FILE) as z:
        assert not any(n.startswith('xl/externalLinks/') for n in z.namelist())
        assert not any(re.search(r'(?<![A-Za-z])(?:[A-Za-z]:[\\/]|file:/)', z.read(n).decode()) for n in z.namelist() if n.endswith(('.xml','.rels')))
        core=ET.fromstring(z.read('docProps/core.xml'))
        assert not any(e.text for e in core if e.tag.rsplit('}',1)[-1] in ('creator','lastModifiedBy'))
        charts=sorted(n for n in z.namelist() if re.fullmatch(r'xl/charts/chart\d+\.xml',n))
        assert len(charts)==5
        for i, name in enumerate(charts):
            root=ET.fromstring(z.read(name));lines=root.findall('.//c:lineChart',ns)
            assert len(lines)==1
            series=lines[0].findall('c:ser',ns)
            assert len(series)==[3,2,2,4,2][i]
            count=[]
            for item in series:
                ref=item.find('c:val/c:numRef/c:f',ns).text
                assert ref.endswith(str(46 if i<2 else 331)), ref
                pts=item.findall('c:val/c:numRef/c:numCache/c:pt/c:v',ns)
                count.append(sum(p.text not in (None,'','#N/A') for p in pts))
            points.append(count)
    expected=[[15,7 if fx else 0,7 if fx else 0],[15,7],[300,300],[300,300,244,300],[300,300]]
    assert points==expected, points
    for name in NAMES:
        assert f[name].freeze_panes is None
    result={'task':'TASK-0036','status':'Review','visible_sheets':NAMES,'charts':5,'series':[3,2,2,4,2],
            'chart_numeric_points':points,'cr_full_rows':[15,15,4999,5000,4999],'cf_level_rows':300,
            'cf_bet_conflict_points':56,'vip_cumulative_resolved':True,'CF_SGD_TO_USD':fx,
            'usd_cost_comparison':'待User确认汇率' if fx is None else '按User汇率换算',
            'intentional_chart_NA':len(expected_na),'unexpected_formula_errors':0,'missing_formula_caches':0,
            'prior_model_values_changed':0,'external_links':0,'frozen_sheets':0}
    f.close();v.close()
    return result


def main() -> None:
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('action',choices=['extract','sanitize','verify'])
    p.add_argument('--cr',type=Path,required=True);p.add_argument('--cf',type=Path,required=True)
    p.add_argument('--output',type=Path,required=True);p.add_argument('--sgd-usd',type=float)
    a=p.parse_args()
    assert not a.output.resolve().is_relative_to(Path(__file__).resolve().parents[4])
    a.output.mkdir(parents=True,exist_ok=True)
    if a.action=='sanitize':
        print(sanitize(a.output,FILE));return
    data=extract(a.cr,a.cf,a.sgd_usd)
    result=data if a.action=='extract' else verify(data,a.output)
    name='comparison-inputs.json' if a.action=='extract' else 'validation.json'
    (a.output/name).write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'action':a.action,'status':'passed','output':name}))


if __name__=='__main__':
    main()
