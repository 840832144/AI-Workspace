"""Top Tycoon 录屏 Spin 样本复算（证据版）。

读取 Top_Tycoon数值模型.xlsx 的 Spin样本 工作表，复算各 video×BET 切片
的概率样本统计，用于校验《Top Tycoon 数值调研报告》中的核心数字。

用法：

    python -X utf8 数值策划/工具/recompute_tt_spin_stats.py \
      --input "数值文档/03_分析与复盘/Top_Tycoon素材/Top_Tycoon数值模型.xlsx"
"""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from pathlib import Path
from typing import DefaultDict, Dict, List, Tuple

from openpyxl import load_workbook

SHEET_SPIN_SAMPLE = "Spin样本"
SAMPLE_MARK = "是"
SINGLE_SPIN = "1"


def _clean(value: object) -> str:
    """清洗导出文本：去掉首尾单引号，避免 '是' 与 '1' 无法匹配。"""
    text = "" if value is None else str(value).strip()
    while len(text) >= 2 and text[0] == "'" and text[-1] == "'":
        text = text[1:-1]
    return text


def _to_float(value: str) -> float:
    try:
        return float(value or 0)
    except ValueError:
        return 0.0


def load_single_spin_samples(path: Path) -> List[Dict[str, str]]:
    """读取概率方向性样本：probability_sample=是 且 spin_count=1。"""
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[SHEET_SPIN_SAMPLE]
    header: List[str] | None = None
    samples: List[Dict[str, str]] = []
    for row in worksheet.iter_rows(values_only=True):
        values = [_clean(value) for value in row]
        if header is None:
            header = values
            continue
        record = dict(zip(header, values))
        if record.get("probability_sample") == SAMPLE_MARK and record.get("spin_count") == SINGLE_SPIN:
            samples.append(record)
    workbook.close()
    return samples


def aggregate(
    samples: List[Dict[str, str]],
) -> Dict[Tuple[str, str], Dict[str, object]]:
    """按 video_id × bet_multiplier 聚合，输出调研报告 §5 所需的切片统计。"""
    buckets: DefaultDict[Tuple[str, str], Dict[str, object]] = defaultdict(
        lambda: {
            "n": 0,
            "coin": 0.0,
            "coin_hit": 0,
            "energy": 0.0,
            "progress": 0.0,
            "types": Counter(),
        }
    )
    for record in samples:
        key = (record["video_id"], record["bet_multiplier"])
        bucket = buckets[key]
        bucket["n"] = int(bucket["n"]) + 1
        coin = _to_float(record["coin_gain_observed"])
        bucket["coin"] = float(bucket["coin"]) + coin
        if coin > 0:
            bucket["coin_hit"] = int(bucket["coin_hit"]) + 1
        bucket["energy"] = float(bucket["energy"]) + _to_float(record["energy_return_observed"])
        bucket["progress"] = float(bucket["progress"]) + _to_float(record["progress_gain_observed"])
        bucket["types"].update([record["primary_category"]])
    return buckets


def print_table(buckets: Dict[Tuple[str, str], Dict[str, object]]) -> None:
    """输出 Markdown 风格结果表。"""
    print("| 视频/下注 | 有效 n | 金币命中率 | 均次金币 | 金币/耗能 | 均次能量返还 | 观察自返还率 | 均次进度 |")
    print("|---|---:|---:|---:|---:|---:|---:|---:|")
    for key in sorted(buckets):
        bucket = buckets[key]
        n = int(bucket["n"])
        bet = float(key[1])
        coin_mean = float(bucket["coin"]) / n
        energy_mean = float(bucket["energy"]) / n
        print(
            f"| {key[0]} x{key[1]} | {n} | "
            f"{float(bucket['coin_hit']) / n * 100:.1f}% | {coin_mean:,.2f} | "
            f"{coin_mean / bet:,.2f} | {energy_mean:.4f} | "
            f"{energy_mean / bet * 100:.1f}% | {float(bucket['progress']) / n:.4f} |"
        )


def print_overall(samples: List[Dict[str, str]]) -> None:
    """输出全样本混合口径：仅用于说明选择偏置，不作总体推断。"""
    cost = sum(_to_float(record["bet_multiplier"]) for record in samples)
    energy = sum(_to_float(record["energy_return_observed"]) for record in samples)
    coin = sum(_to_float(record["coin_gain_observed"]) for record in samples)
    progress = sum(_to_float(record["progress_gain_observed"]) for record in samples)
    print()
    print(
        f"全样本混合口径：{len(samples)} 行，消耗能量 {cost:,.0f}，返还能量 {energy:,.1f}"
        f"（观察自返还率 {energy / cost * 100:.1f}%），金币 {coin:,.0f}，进度 {progress:,.0f}"
    )
    print("注意：混合口径受片段选择与离散大额返还影响，只用于说明样本偏置，不作 RTP 推断。")


def main() -> None:
    parser = argparse.ArgumentParser(description="Top Tycoon Spin 样本复算")
    parser.add_argument("--input", required=True, type=Path, help="Top_Tycoon数值模型.xlsx 路径")
    args = parser.parse_args()

    samples = load_single_spin_samples(args.input)
    if not samples:
        raise SystemExit("未找到符合条件（probability_sample=是 且 spin_count=1）的样本")
    print(f"有效单 Spin 样本：{len(samples)} 行")
    print()
    print_table(aggregate(samples))
    print_overall(samples)


if __name__ == "__main__":
    main()
