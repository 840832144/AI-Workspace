"""Prepare/verify controlled three-Bet comparison. Sources are read-only.

prepare --baseline <threeway1500> --current <source-lock directory>
        --pop-source <PS_HANDOFF> --out <controlled output>
cache-charts / verify --out <controlled output>
No collector, game configuration, SVN write, or online analytics operation.
"""
from __future__ import annotations

import argparse
import bisect
import csv
import json
import math
import re
import subprocess
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZipFile

import openpyxl
from build_cr_cf_1500 import read, save, cache_charts
from build_cashfrenzy_curves import sanitize


def prepare(baseline: Path, current: Path, pop_source: Path, out: Path) -> dict:
    out.mkdir(parents=True, exist_ok=True)
    data=read(baseline/'inputs.controlled.json')
    old=read(baseline/'expected.controlled.json')
    raw=read(pop_source/'PS_DATASET.json')
    tables=read(current/'cr-current-tables.controlled.json')
    lock=read(current/'source-lock.local.json')
    def svn(*args: str) -> ET.Element:
        r=subprocess.run(['svn',*args],capture_output=True,encoding='utf-8')
        if r.returncode: raise RuntimeError('SVN read failed; no configuration written')
        return ET.fromstring(r.stdout)
    revision=int(svn('info','--xml',lock['url']).find('entry').attrib['revision'])
    paths=[]
    if revision>lock['revision']:
        log=svn('log','--xml','-v','-r',f"{lock['revision']+1}:{revision}",lock['url'])
        paths=[p.text for p in log.findall('.//path') if Path(p.text).stem in tables]
    if paths: raise RuntimeError('Relevant trunk changes: refresh a consistent source before preparing')
    # Reuse unchanged source values; record applicability at the checked HEAD.
    save(out/'source-lock.local.json',{'revision':revision,'prior_revision':lock['revision'],
         'read_time_utc':datetime.now(timezone.utc).isoformat(),'related_paths_changed':0,
         'source_write':False,'hash_checks':False})

    # Targeted evidence across all three sessions: only level events and store timeline.
    level_events=[];counts={}
    for name in ['ps_run1.jsonl','ps_run2.jsonl','ps_run3.jsonl']:
        count=0
        for line in (pop_source/'raw'/name).open(encoding='utf-8'):
            r=json.loads(line)
            if r.get('kind') not in ('event','ev','levelup-data'): continue
            if r.get('kind')!='levelup-data' and (r.get('name') or r.get('s0'))!='LevelProgressUpdateEvent':continue
            j=json.loads(r.get('json') or '{}')
            if isinstance(j,dict) and j.get('level'):
                level_events.append((float(r['ts']),int(j['level'])));count+=1
        counts[name]=count
    level_events.sort();times=[x[0] for x in level_events]
    hist=list(csv.DictReader((pop_source/'dataset/ps_store_history.csv').open(encoding='utf-8-sig',newline='')))
    timeline=[]
    for r in hist:
        ts=float(r['first_ts']);i=bisect.bisect_right(times,ts)-1
        lv=level_events[i][1] if i>=0 else 1
        timeline.append([lv,int(re.search(r'M(\d+)$',r['offer_id'])[1]),int(r['usd99_amount']),ts])
    changes=[]
    for lv,m,_,_ in timeline:
        if not changes or changes[-1][1]!=m:changes.append([lv,m])
    assert changes==[[20,200000],[50,2200000],[75,3600000]],('Multiplier alignment requires review',changes)
    # First observed store multiplier is L20. Earlier levels use that baseline
    # as an explicit scenario; never relabel the first observation as L1.
    observed_changes=changes
    changes=[[1,changes[0][1]],*changes[1:]]
    levels=list(csv.DictReader((pop_source/'dataset/ps_levels_full.csv').open(encoding='utf-8-sig',newline='')))
    assert [int(r['level']) for r in levels]==list(range(1,89))
    pop_levels=[[int(r['level']),int(r['start_xp']),int(r['next_xp']),int(r['span_xp'])] for r in levels]
    assert all(r[2]-r[1]==r[3] for r in pop_levels)
    assert all(a[2]==b[1] for a,b in zip(pop_levels,pop_levels[1:]))
    assert pop_levels==[[r['level'],r['start_xp'],r['next_xp'],r['span_xp']] for r in raw['level_table']]
    cap=[[r['level'],r['max_bet_per_line']] for r in raw['recommended_bet']['level_cap_ladder_per_line']]
    # Source-defined initial and machine-specific known bounds; not observed played-bet extrema.
    pop_min=min(r[4] for r in data['pop']['levels'])
    cf_min=min(r[1] for r in data['cf'])
    bids={r['levelId']:r for r in tables['SlotsCasinoBetList']['records']}
    pools={}
    for r in tables['SlotsCasinoBetUnlock']['records']:
        if r['highroller']==0:pools.setdefault(r['level'],[]).append(bids[r['betlevel']])
    assert all(r['bet2']==r['levelExp']*10000 for r in bids.values())
    assert next(r['val'] for r in tables['CommCfg']['records'] if r['id']==161)==1
    xp_anchor=pop_levels[-1][3]
    slope=sum((r[0]-88)*(r[3]-xp_anchor) for r in pop_levels[74:])/sum((r[0]-88)**2 for r in pop_levels[74:])
    assert slope>0
    rows=[];expected=[];stages={1,5,50,75,89,1501};last=None
    for i,(cr,cf,prev) in enumerate(zip(data['cr'],data['cf'],old)):
        lv=i+1;stage=max(k for k in pools if k<=lv);pool=sorted(pools[stage],key=lambda r:r['bet2'])
        minimum=pool[0]['bet2'];maximum=pool[-1]['bet2']
        per_line=next((v for l,v in reversed(cap) if l<=lv),pop_min/20)
        m=next(v for l,v in reversed(changes) if l<=lv)
        need=pop_levels[i][3] if lv<=88 else math.floor((xp_anchor+slope*(lv-88))/1e6+0.5)*1e6
        source=[lv,cr[1],cr[2],minimum,maximum,cr[5],cf[1],prev[9],prev[8],
                pop_levels[i][3] if lv<=88 else None,per_line,m,stage]
        rows.append(source)
        rates=[cr[5],500000*prev[8],10*m/1.99]
        mins=[minimum,cf_min,pop_min];maxs=[maximum,cf[1],per_line*20]
        needs=[cr[2],prev[9],need];units=[1/10000,1/3,1]
        models=[]
        for g in range(3):
            low,high=mins[g],maxs[g];rec=max(low,min(high,rates[g]*100/37))
            bets=[low,rec,high];spins=[needs[g] if g==0 and cr[1]==1 else needs[g]/(b*units[g]) for b in bets]
            costs=[n*b/rates[g] for n,b in zip(spins,bets)]
            models.append([lv,*bets,*[b/rates[g] for b in bets],*spins,*costs,rates[g],needs[g]])
        target=pop_min/rates[2]*rates[0]
        eligible=[r for r in pool if r['bet2']>=target]
        chosen=minimum if lv<5 else (eligible[0]['bet2'] if eligible else maximum)
        expected.append({'models':models,'floor':chosen,'target':target,'stage':stage})
        sig=(stage,cr[5],cf[1],prev[8],per_line,m)
        if last is not None and sig!=last:stages.add(lv)
        last=sig
    vip=[]
    new={r['tier']:r for r in raw['vip']['ladder']}
    for i,r in enumerate(data['pop']['vip']):
        t=i+1
        if t in new:vip.append([t,new[t]['tp_to_obtain'],new[t]['chip_package_bonus'],'10/10界面'])
        elif t<=3:vip.append([t,r[1],r[2],'10/09界面'])
        else:vip.append([t,r[1],None,'门槛历史；权益拟合'])
    data.update(book=f'CR_CF_POP_三档Bet与消耗对比_trunk_r{revision}_L88.xlsx',revision=revision,
                level_inputs=rows,pop88=pop_levels,pop_vip=vip,pop_multiplier=changes,
                stages=[[a,b-1] for a,b in zip(sorted(stages),sorted(stages)[1:])],
                bet_pool=[[stage,r['levelId'],r['bet2'],r['levelExp']] for stage,rs in sorted(pools.items()) for r in rs],
                controls={'balance_usd':100,'recommended_divisor':37,'pop_low_price':1.99,
                          'pop_low_k':10,'pop_base_m':200000,'pop_campaign':0,'pop_vip_bonus':0,
                          'cf_base_rate':500000,'pop_lines':20,'xp_anchor_level':88,'xp_slope':slope,
                          'xp_round':1000000,'pop_min':pop_min,'cf_min':cf_min},
                observed_recommended=raw['recommended_bet']['observed'],store_timeline=timeline,
                multiplier_first_observed=observed_changes,
                raw_read_scope={'files':counts,'scope':'level updates only; store-history timestamps aligned'})
    save(out/'inputs.controlled.json',data);save(out/'bet-expected.controlled.json',expected)
    save(out/'preparation-validation.json',{'revision':revision,'reused_CR_tables':7,'related_changes':0,
         'POP_source_rows':88,'POP_fit_rows':1412,'multiplier_changes':[r[0] for r in changes],
         'stage_count':len(data['stages']),'source_write':False,'full_raw_analysis':False})
    return {'revision':revision,'stages':len(data['stages']),'raw_level_events':sum(counts.values()),'ready':True}


def verify(out: Path) -> dict:
    d=read(out/'inputs.controlled.json');book=d['book'];sanitize(out,book)
    wb=openpyxl.load_workbook(out/book,data_only=True)
    form=openpyxl.load_workbook(out/book,data_only=False)
    expected=read(out/'bet-expected.controlled.json')
    for i,row in enumerate(expected):
        for g,name in enumerate(['CR明细','CF明细','POP明细']):
            actual=[wb[name].cell(i+6,c).value for c in range(1,16)]
            assert all(isinstance(a,(float,int)) and math.isclose(a,b,rel_tol=1e-9,abs_tol=1e-6) for a,b in zip(actual,row['models'][g])),(name,i+1,actual,row['models'][g])
        assert wb['CR调整候选'].cell(i+6,5).value==row['floor'],('candidate',i+1)
    # Candidate is expressed entirely in the current ordinary-machine pool.
    pools={}
    for stage,bid,bet,exp in d['bet_pool']:pools.setdefault(stage,[]).append((bid,bet))
    candidates=[];previous=0
    for i,row in enumerate(expected):
        r=i+6;c=wb['CR调整候选'];floor=c.cell(r,14).value
        assert floor>=previous and floor in [bet for _,bet in pools[row['stage']]]
        assert floor<=row['models'][0][3]
        if i>=4:
            assert math.isclose(c.cell(r,12).value,1,rel_tol=1e-9)
            assert all(math.isclose(m[10],m[11],rel_tol=1e-9) and math.isclose(m[11],m[12],rel_tol=1e-9) for m in row['models'])
        previous=floor
        for label,value in [('美元对齐',c.cell(r,5).value),('不回退',floor)]:
            keep=[bid for bid,bet in pools[row['stage']] if bet>=value]
            drop=[bid for bid,bet in pools[row['stage']] if bet<value]
            candidates.append([label,i+1,row['stage'],value,','.join(map(str,keep)),','.join(map(str,drop))])
    with (out/'CR_Bet最低档候选_逐级.csv').open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f);w.writerow(['方案','等级','当前解锁池阶段','候选最低Bet','保留BetID','限制BetID']);w.writerows(candidates)
    errors=[(s.title,c.coordinate,c.value) for s in wb for row in s for c in row if c.data_type=='e']
    assert not errors,errors[:10]
    assert all(s.freeze_panes is None for s in form)
    charts=sum(len(s._charts) for s in form)
    assert charts==7,charts
    with ZipFile(out/book) as z:
        assert not any(n.startswith('xl/externalLinks/') for n in z.namelist())
        ns={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart'};points=[]
        for name in z.namelist():
            if '/charts/' not in name or not name.endswith('.xml'):continue
            root=ET.fromstring(z.read(name));sizes=[]
            for ref in root.findall('.//c:numRef',ns):
                sn,area=ref.find('c:f',ns).text.rsplit('!',1);sn=sn.strip("'")
                a,b,c,e=openpyxl.utils.range_boundaries(area)
                values=[wb[sn].cell(r,k).value for r in range(b,e+1) for k in range(a,c+1)]
                pts=ref.findall('c:numCache/c:pt',ns)
                assert len(pts)==len(values)
                assert all(math.isclose(float(p.find('c:v',ns).text),values[int(p.attrib['idx'])],rel_tol=1e-12,abs_tol=1e-8) for p in pts)
                sizes.append(len(pts))
            points.append(sizes)
    result={'levels':1500,'games':3,'Bet_cases':3,'POP_source_rows':88,'POP_fit_rows':1412,
            'stage_count':len(d['stages']),'charts':charts,'chart_points':points,'formula_errors':0,
            'external_links':0,'freeze_panes':0,'native_WPS_open':False,'source_write':False,
            'recommended_is_scenario':True,'POP_cap_20_lines_is_conditional':True,
            'candidate_pool_membership':True,'no_rollback_floor_monotone':True,
            'EXP_mode_level_gross_invariance':True}
    samples=[]
    for lv in [1,5,10,20,25,30,40,49,50,75,88,100,300,500,1000,1500]:
        r=lv+5;c=wb['CR调整候选'];models=expected[lv-1]['models']
        samples.append([lv,*[m[4] for m in models],c.cell(r,5).value,c.cell(r,14).value,*[m[10] for m in models]])
    save(out/'key-samples.controlled.json',samples)
    lines=['# CR / CF / POP 三档Bet与消耗比较（受控）','',
      '当前CR trunk r'+str(d['revision'])+'；POP L88新包；CF复用既有模型。名义100美元余额场景，可在Excel首页B11修改。',
      '', '## 先看结论', '',
      '- POP底档美元成本并非始终高于另外两款；最低档及兑换情景需同时看。',
      '- POP商城M最早在L20观测，L50和L75分别出现变化；L1–19回推、L89+延续末档是情景。',
      '- 想提高每转消耗，优先审最低Bet和推荐，而不是只提高最大Bet。CR候选使用当前已存在的普通机台档位。',
      '- 想提高升一级总消耗，单独提高Bet且让EXP同比增长不够；当前CR 5级起理论毛下注总成本不变。',
      '', '## 关键节点（USD为标准低档、VIP0、无促销的毛下注价值）','',
      '|等级|CR最低每转$|CF底档每转$|POP底档每转$|CR美元对齐Bet|CR不回退Bet|CR升级毛$|CF升级毛$|POP升级毛$|',
      '|---:|---:|---:|---:|---:|---:|---:|---:|---:|']
    for row in samples:lines.append('|'+ '|'.join(str(v) if j==0 else f'{v:,.0f}' if j in (4,5) else f'{v:,.4f}' for j,v in enumerate(row))+'|')
    lines += ['', '## 建议如何落到Bet配置', '',
      '1. 保留1–4级。5级起按POP底档美元价值换算CR金币目标，再向上取当前解锁池已有档；先看现档跳幅是否可接受。',
      '2. 若不希望升到50/75级时最低Bet突然回退，使用不回退候选；它会高于POP的同价值目标，差额见工作簿。',
      '3. 配置落点是普通模式SlotsCasinoBetUnlock的可选池限制，按新边界复制当前适用阶段并保留所有非Bet字段；BetList数值/EXP、BetShow映射、机台换算系数及HighRoller池不变。普通机台共用池会影响其可选档，不能说其他机台体验完全不变。CSV只是逐级方案，不可直接导入。',
      '4. 推荐Bet目前是余额除数近似情景，不把37直接写进CommCfg。两个真实样本并不能证明精确取整或筛档算法；实际推荐还需独立核对。',
      '5. 最大Bet维持CR当前值。POP新上限字段每线/总注未完全闭合，不能据此直接上调最大Bet。',
      '6. 如目标改为每级更贵，需要另审每Spin经验与Bet或等级门槛的关系；这会改变已确认升级体验，本轮没有执行。',
      '', '## 证据边界', '',
      '- POP最低档只在早期20线机台有参考，高级别真实下限未确认；CF同样只是已知最低档延续。',
      '- POP最大Bet以20线条件计算，可改首页线数看50线敏感性；L1–3基础档不是独立证实的最大档。',
      '- POP 1级门槛是交接方回填，2–88保留已知跨度，89–1500按75–88趋势拟合；不改87级原有下降。',
      '- POP付费Spin EXP=Bet沿用提供方模型，逐局Bet由经验差反推不能自证该规律；免费转除外。',
      '- 升级成本是理论毛下注，整局取整另列；未把RTP样本、等级奖励或真实付款混入，也不能直接据此判断缺币概率。',
      '- CF实测/历史/拟合边界留在每行来源；没有重做TASK-0033/0034或新增采集。',
      '', '## 验证与交付', '',
      '工作簿12页、7张原生折线图、99个等级段（每段三方各一行）、每款1500行明细；全部对应公式与独立计算。0公式错误/外链/冻结；WPS原生打开未执行。',
      '源数据、CR配置和SVN没有写入。CSV是受控提案，PR #10等待Review，不合并/finalize。']
    (out/'阅读与调整建议.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
    save(out/'validation-summary.json',result);wb.close();form.close();return result


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('action',choices=['prepare','cache-charts','verify'])
    for n in ['baseline','current','pop-source','out']:p.add_argument('--'+n,type=Path,required=n=='out')
    a=p.parse_args()
    r=prepare(a.baseline,a.current,a.pop_source,a.out) if a.action=='prepare' else cache_charts(a.out) if a.action=='cache-charts' else verify(a.out)
    print(json.dumps(r,ensure_ascii=False))
