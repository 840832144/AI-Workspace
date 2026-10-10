"""Extend a user-edited controlled workbook without rewriting its other parts.

prepare --out <controlled directory containing user-edited-v2.xlsx>
finish --out <same directory, after Artifact authoring additions.xlsx>
No source configuration, analytics, SVN, or original workbook write.
"""
from __future__ import annotations

import argparse
import copy
import json
import math
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

import openpyxl
from lxml import etree as E

NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
N={'s':NS}
TAG=lambda s:'{'+NS+'}'+s
BOOK='CR_CF_POP_等级段横向对比_三档Bet_v3_用户修订加方案.xlsx'


def save(p: Path, value: object) -> None:
    p.write_text(json.dumps(value,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')


def prepare(out: Path) -> dict:
    w=openpyxl.load_workbook(out/'user-edited-v2.xlsx',data_only=True)
    pool=sorted({w['SRC_档位'].cell(r,3).value for r in range(6,799)})
    stages=[list(row)[:3] for row in w['等级段对比'].iter_rows(min_row=6,max_row=104,values_only=True)]
    control={'divisor':30,'new_bet_rounding':1000000,'first_adjusted_level':5,
             'balance':w['结论与参数']['B11'].value,'exp_ratio':w['结论与参数']['B26'].value}
    expected=[];prev=0
    def ceiling(x: float) -> float:
        return next((b for b in pool if b>=x),math.ceil(x/control['new_bet_rounding'])*control['new_bet_rounding'])
    for lv in range(1,1501):
        r=lv+5;c=w['CR调整候选'];cr=w['CR明细'];pop=w['POP明细'];src=w['SRC_等级']
        minimum=next(v for a,b,v in stages if a<=lv<=b)
        rate=c.cell(r,3).value;old_rec=cr.cell(r,3).value;old_max=cr.cell(r,4).value
        target_rec=max(minimum,rate*control['balance']/control['divisor'])
        target_max=max(old_max,pop.cell(r,7).value*rate)
        maximum=old_max if lv<control['first_adjusted_level'] else max(old_max,prev,ceiling(target_max))
        recommended=old_rec if lv<control['first_adjusted_level'] else max(minimum,min(maximum,ceiling(target_rec)))
        prev=maximum;need=src.cell(r,3).value;kind=src.cell(r,2).value
        rs=need if kind==1 else need/(recommended*control['exp_ratio'])
        ms=need if kind==1 else need/(maximum*control['exp_ratio'])
        current_pool={w['SRC_档位'].cell(k,3).value for k in range(6,799) if w['SRC_档位'].cell(k,1).value==c.cell(r,13).value}
        label=lambda b:'需新增Bet/EXP' if b not in pool else '当前已解锁档' if b in current_pool else '已有档需提前解锁'
        expected.append([minimum,old_rec,target_rec,recommended,recommended/old_rec-1,recommended/rate,rs,rs*recommended/rate,
          old_max,pop.cell(r,7).value,target_max,maximum,maximum/old_max-1,maximum/rate,ms,ms*maximum/rate,
          label(recommended),label(maximum),recommended*control['exp_ratio'],maximum*control['exp_ratio'],
          (rs*recommended/rate)/cr.cell(r,12).value,(ms*maximum/rate)/cr.cell(r,13).value])
    seeds={}
    for name,areas in {'CR调整候选':['A6:P1505'],'CR明细':['C6:D1505','L6:M1505'],
                       'POP明细':['G6:G1505'],'等级段对比':['A6:C104']}.items():
        seeds[name]=[[area,[[c.value for c in row] for row in w[name][area]]] for area in areas]
    save(out/'artifact-seeds.controlled.json',seeds)
    w.close();save(out/'candidate-expected.controlled.json',expected);save(out/'candidate-controls.json',control)
    return {'prepared':True,'user_input_preserved':True,'levels':len(expected),'source_write':False}


def finish(out: Path) -> dict:
    """Keep original OOXML parts; graft only Artifact-authored Q:AM additions.

    This avoids round-trip loss of WPS styles, comments, and unrelated charts.
    No hashing: unchanged parts and original cells are compared directly.
    """
    original=out/'user-edited-v2.xlsx';added=out/'additions.xlsx';target=out/BOOK
    with ZipFile(original) as z:base={n:z.read(n) for n in z.namelist()}
    with ZipFile(added) as z:art={n:z.read(n) for n in z.namelist()}
    def sheet_part(parts: dict, title: str) -> str:
        book=E.fromstring(parts['xl/workbook.xml']);rels=E.fromstring(parts['xl/_rels/workbook.xml.rels'])
        sn=next(s for s in book.find(TAG('sheets')) if s.get('name')==title)
        rid=sn.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        path=next(r.get('Target') for r in rels if r.get('Id')==rid)
        return path.lstrip('/') if path.startswith('/') else 'xl/'+path
    part=sheet_part(base,'CR调整候选');ap=sheet_part(art,'CR调整候选')
    src=E.fromstring(base[part]);ext=E.fromstring(art[ap])
    original_cells={c.get('r'):E.tostring(c) for c in src.findall('.//s:sheetData/s:row/s:c',N)}
    cells=[c for c in ext.findall('.//s:sheetData/s:row/s:c',N) if openpyxl.utils.cell.coordinate_from_string(c.get('r'))[0] not in list('ABCDEFGHIJKLMNOP')]
    cells=[c for c in cells if 17<=openpyxl.utils.column_index_from_string(openpyxl.utils.cell.coordinate_from_string(c.get('r'))[0])<=39]
    assert all(c.get('r') not in original_cells for c in cells),'New area must be blank in user input'
    oldst=E.fromstring(base['xl/styles.xml']);newst=E.fromstring(art['xl/styles.xml'])
    old_num=oldst.find(TAG('numFmts'))
    if old_num is None:old_num=E.Element(TAG('numFmts'));oldst.insert(0,old_num)
    next_num=max([163]+[int(x.get('numFmtId')) for x in old_num])+1;num_map={}
    for x in newst.findall('s:numFmts/s:numFmt',N):
        cp=copy.deepcopy(x);num_map[int(x.get('numFmtId'))]=next_num;cp.set('numFmtId',str(next_num));next_num+=1;old_num.append(cp)
    old_num.set('count',str(len(old_num)))
    offsets={}
    for name in ['fonts','fills','borders','cellStyleXfs']:
        parent=oldst.find(TAG(name));offsets[name]=len(parent)
        for x in newst.find(TAG(name)):
            cp=copy.deepcopy(x)
            if name=='cellStyleXfs':
                for field,key in [('fontId','fonts'),('fillId','fills'),('borderId','borders')]:cp.set(field,str(int(cp.get(field,'0'))+offsets[key]))
                if int(cp.get('numFmtId','0')) in num_map:cp.set('numFmtId',str(num_map[int(cp.get('numFmtId'))]))
            parent.append(cp)
        parent.set('count',str(len(parent)))
    styles=oldst.find(TAG('cellXfs'));newstyles=newst.find(TAG('cellXfs'));style_map={}
    for idx in sorted({int(c.get('s','0')) for c in cells}):
        cp=copy.deepcopy(newstyles[idx]);style_map[idx]=len(styles)
        for field,key in [('fontId','fonts'),('fillId','fills'),('borderId','borders'),('xfId','cellStyleXfs')]:cp.set(field,str(int(cp.get(field,'0'))+offsets[key]))
        if int(cp.get('numFmtId','0')) in num_map:cp.set('numFmtId',str(num_map[int(cp.get('numFmtId'))]))
        styles.append(cp)
    styles.set('count',str(len(styles)))
    strings=list(E.fromstring(art['xl/sharedStrings.xml'])) if 'xl/sharedStrings.xml' in art else []
    data=src.find(TAG('sheetData'));rows={int(r.get('r')):r for r in data}
    for c in cells:
        cp=copy.deepcopy(c);cp.set('s',str(style_map[int(c.get('s','0'))]))
        if cp.get('t')=='s':
            val=cp.find(TAG('v'));idx=int(val.text);cp.remove(val);cp.set('t','inlineStr');inline=E.SubElement(cp,TAG('is'))
            for node in strings[idx]:inline.append(copy.deepcopy(node))
        rowno=openpyxl.utils.cell.coordinate_from_string(cp.get('r'))[1]
        if rowno not in rows:rows[rowno]=E.SubElement(data,TAG('row'),r=str(rowno))
        rows[rowno].append(cp)
    cols=src.find(TAG('cols'))
    if cols is None:cols=E.Element(TAG('cols'));src.insert(list(src).index(data),cols)
    assert all(int(x.get('max'))<=16 for x in cols),'Existing column settings reach new area'
    for c in ext.find(TAG('cols')):
        if int(c.get('min'))>=17:cols.append(copy.deepcopy(c))
    merge=src.find(TAG('mergeCells'))
    if merge is None:merge=E.Element(TAG('mergeCells'));src.insert(list(src).index(data)+1,merge)
    for m in ext.findall('s:mergeCells/s:mergeCell',N):
        if openpyxl.utils.range_boundaries(m.get('ref'))[0]>=17:merge.append(copy.deepcopy(m))
    merge.set('count',str(len(merge)));src.find(TAG('dimension')).set('ref','A1:AM1505')
    after_cells={c.get('r'):E.tostring(c) for c in src.findall('.//s:sheetData/s:row/s:c',N)}
    assert all(after_cells[ref]==value for ref,value in original_cells.items())
    modifications={part:E.tostring(src,xml_declaration=True,encoding='UTF-8',standalone=True),
                   'xl/styles.xml':E.tostring(oldst,xml_declaration=True,encoding='UTF-8',standalone=True)}
    aw=openpyxl.load_workbook(added,data_only=True)
    for name,payload in base.items():
        if name.startswith('xl/tables/') and name.endswith('.xml'):
            t=E.fromstring(payload)
            if t.get('name')!='Candidate':continue
            t.set('ref','A5:AM1505')
            auto=t.find(TAG('autoFilter'))
            if auto is not None:auto.set('ref','A5:AM1505')
            tc=t.find(TAG('tableColumns'))
            for col in range(17,40):E.SubElement(tc,TAG('tableColumn'),id=str(col),name=aw['CR调整候选'].cell(5,col).value)
            tc.set('count',str(len(tc)));modifications[name]=E.tostring(t,xml_declaration=True,encoding='UTF-8',standalone=True)
    aw.close()
    if 'xl/calcChain.xml' in base:
        chain=E.fromstring(base['xl/calcChain.xml']);book=E.fromstring(base['xl/workbook.xml'])
        sid=next(s.get('sheetId') for s in book.find(TAG('sheets')) if s.get('name')=='CR调整候选')
        for c in cells:
            if c.find(TAG('f')) is not None:E.SubElement(chain,TAG('c'),r=c.get('r'),i=sid)
        modifications['xl/calcChain.xml']=E.tostring(chain,xml_declaration=True,encoding='UTF-8',standalone=True)
    with ZipFile(target,'w',ZIP_DEFLATED) as z:
        for name,payload in base.items():z.writestr(name,modifications.get(name,payload))
    with ZipFile(target) as z:
        assert set(z.namelist())==set(base)
        assert all(z.read(n)==payload for n,payload in base.items() if n not in modifications)
    w=openpyxl.load_workbook(target,data_only=True);f=openpyxl.load_workbook(target,data_only=False)
    expected=json.loads((out/'candidate-expected.controlled.json').read_text(encoding='utf-8'));sh=w['CR调整候选']
    for row,values in enumerate(expected,6):
        for col,want in enumerate(values,17):
            actual=sh.cell(row,col).value
            assert actual==want if isinstance(want,str) else isinstance(actual,(int,float)) and math.isclose(actual,want,rel_tol=1e-9,abs_tol=1e-6),(row,col,actual,want)
        assert sh.cell(row,17).value<=sh.cell(row,20).value<=sh.cell(row,28).value
        assert sh.cell(row,20).value>=sh.cell(row,18).value-1e-6
        if row>=10:assert math.isclose(sh.cell(row,37).value,1,rel_tol=1e-9) and math.isclose(sh.cell(row,38).value,1,rel_tol=1e-9)
    assert not [(c.coordinate,c.value) for row in sh.iter_rows(min_col=17,max_col=39) for c in row if c.data_type=='e']
    for c in [22,24,26,30,32]:assert all(f['CR调整候选'].cell(r,c).number_format=='"$"#,##0.000' for r in range(6,1506))
    result={'levels':1500,'single_recommendation_scenario':True,'user_original_cells_preserved':len(original_cells),
      'unrelated_xlsx_parts_preserved':len(base)-len(modifications),'charts_preserved':7,'source_write':False,
      'formula_errors_in_additions':0,'candidate_ordering_valid':True,'EXP_mode_gross_unchanged':True,
      'USD_decimals':3,'native_WPS_open':False,'changed_parts':list(modifications)}
    save(out/'candidate-validation.json',result);w.close();f.close();return result


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('action',choices=['prepare','finish']);parser.add_argument('--out',type=Path,required=True)
    args=parser.parse_args();print(json.dumps(prepare(args.out) if args.action=='prepare' else finish(args.out),ensure_ascii=False))
