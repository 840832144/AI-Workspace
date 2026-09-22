"""TASK-0036 current-CF controlled candidates. No SVN writes or source edits.

prepare|verify --base <controlled numerics root> --out <controlled output>
Native authoring uses existing proven exact-edit fallback for imported config formulas.
Report authoring: render_cr_cf_alignment.mjs (Artifact). openpyxl is read-only.
"""
from __future__ import annotations
import argparse
import csv
import json
import math
import subprocess
import xml.etree.ElementTree as ET
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
import openpyxl
from producer_workbook_sources import read_source
from build_cr_bet_dev import grid
from build_cashfrenzy_curves import sanitize

FILES=['LevelCfg','SlotsCasinoBetList','SlotsCasinoBetUnlock','SlotsCasinoBetShow','PriceCheatSheet']
BOOK='CR_CF_等级对齐候选_5000级.xlsx'


def rounded(x: float) -> int:
    return int(Decimal(str(x)).quantize(Decimal('1'),rounding=ROUND_HALF_UP))


def save(path: Path, data: object) -> None:
    path.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')


def prepare(base: Path, out: Path) -> dict:
    source=out/'source'
    receipt=json.loads((out/'source.local.json').read_text(encoding='utf-8'))
    assert receipt['environment']=='CR dev'
    tables={name:read_source(source/(name+'.xlsx')) for name in FILES}
    live=json.loads((base/'outputs/task0036-cf-live-20260922/live-inputs.json').read_text(encoding='utf-8'))
    history=json.loads((base/'outputs/task0036-bet-dev-20260921/bet-dev-inputs.json').read_text(encoding='utf-8'))
    cf=live['cf'];levels=tables['LevelCfg']['records'];bets=tables['SlotsCasinoBetList']['records']
    unlock=tables['SlotsCasinoBetUnlock']['records'];prices=tables['PriceCheatSheet']['records']
    assert len(levels)==5000 and len(bets)==38
    # Smallest positive integer Spin EXP for every existing Bet: gcd of coin stakes.
    # CF raw EXP=coins/3, hence CR EXP=CF EXP*3/coin_unit (not an arbitrary decimal divisor).
    new_coin={r['levelId']:r['bet2'] for r in bets}
    tier_alignment={8:10,10:15,16:30,17:35}  # Stable CR tier ID -> measured CF level.
    new_coin.update({ident:cf[level-1][1] for ident,level in tier_alignment.items()})
    coin_unit=math.gcd(*new_coin.values())
    scale=coin_unit/3
    anchor=cf[49][3];window=cf[29:50]
    slope=sum((r[0]-50)*(r[3]-anchor) for r in window)/sum((r[0]-50)**2 for r in window)
    bet_window=[r for r in cf[29:50] if r[0] in [30,35,40,50]]
    bet_anchor=cf[49][1]
    bet_slope=sum((r[0]-50)*(r[1]-bet_anchor) for r in bet_window)/sum((r[0]-50)**2 for r in bet_window)
    assert slope>0 and bet_slope>0
    changes={n:[] for n in FILES}
    def edit(name: str, row: int, column: str, before: object, after: object, why: str) -> None:
        if before != after:
            changes[name].append([tables[name]['sheet'],f'{column}{row}',before,after,why])
    # Retain stable IDs, align four base-coin tiers. Companion bet multipliers remain original formulas.
    new_exp={k:v//coin_unit for k,v in new_coin.items()}
    assert min(new_exp.values())==1 and all(v%coin_unit==0 for v in new_coin.values())
    assert all(a<b for a,b in zip(new_coin.values(),list(new_coin.values())[1:]))
    assert all(a<=b for a,b in zip(new_exp.values(),list(new_exp.values())[1:]))
    for r in bets:
        ident=r['levelId'];rr=r['_excel_row']
        edit('SlotsCasinoBetList',rr,'F',r['bet2'],new_coin[ident],'最新CF最大Bet必要档位')
        edit('SlotsCasinoBetList',rr,'E',r['levelExp'],new_exp[ident],'按最小Bet公约数换算，1万金币=1EXP')
    for r in tables['SlotsCasinoBetShow']['records']:
        edit('SlotsCasinoBetShow',r['_excel_row'],'D',r['desc'],new_coin[r['betIndex']],'同步变更Bet备注；index/value关系保留')
    ordinary=[r for r in unlock if r['highroller']==0]
    def original_pool(level: int) -> list[dict]:
        stage=max(r['level'] for r in ordinary if r['level']<=level)
        return [r for r in ordinary if r['level']==stage]
    old_coin={r['levelId']:r['bet2'] for r in bets};old_exp={r['levelId']:r['levelExp'] for r in bets}
    def old_id(level: int) -> int:
        return max(original_pool(level),key=lambda r:old_coin[r['betlevel']])['betlevel']
    known={5:cf[5][1]}  # L5 amount inferred from the unchanged server interval through L6.
    for row in cf[5:50]:
        if row[1]!=list(known.values())[-1]:known[row[0]]=row[1]
    cap=[]
    for lv in range(1,5001):
        if lv<5: value=old_coin[old_id(lv)]
        elif lv<75: value=known[max(x for x in known if x<=lv)]
        else:
            stage=lv//25*25
            fit=bet_anchor+bet_slope*(stage-50)
            value=max(v for v in new_coin.values() if v<=fit)
        cap.append(value)
    cap_id={v:k for k,v in new_coin.items()}
    stage_set={r['level'] for r in ordinary}|set(known)|{i+1 for i in range(1,5000) if cap[i]!=cap[i-1]}
    fields=[k for _,k in tables['SlotsCasinoBetUnlock']['fields']]
    ug=grid(source/'SlotsCasinoBetUnlock.xlsx')[tables['SlotsCasinoBetUnlock']['sheet']]
    output=[];unlock_notes=[]
    for stage in sorted(stage_set):
        pool=original_pool(stage);template=max(pool,key=lambda r:old_coin[r['betlevel']])
        keep=[dict(r,level=stage) for r in pool if new_coin[r['betlevel']]<=cap[stage-1]]
        ident=cap_id[cap[stage-1]]
        if ident not in [r['betlevel'] for r in keep]:
            keep.append(dict(template,level=stage,betlevel=ident))
        for r in sorted(keep,key=lambda x:x['betlevel']):
            output.append([r.get(k) for k in fields])
        unlock_notes.append([stage,template['_excel_row'],cap[stage-1]])
    output.extend([[r.get(k) for k in fields] for r in unlock if r['highroller']==1])
    output=[r+[None]*(len(ug[0])-len(r)) for r in output]
    assert len({(r[0],r[1],r[2]) for r in output})==len(output)
    # Current authoritative price table only. Preserve SKU shape, VIP ratios, non-coin currencies.
    shop=sorted([r for r in prices if (r['priceType'],r['vipType'],r['level'])==(9,1,0)],key=lambda r:r['money'])
    basis={r['level']:r for r in prices if (r['priceType'],r['vipType'],r['money'])==(17,1,100)}
    original_base=basis[0]['vip_16_0'];shape={r['money']:r['vip_16_0']/((r['money']+1)/100)/original_base for r in shop}
    assert len(shape)==30
    low,high=min(shape.values()),max(shape.values())
    new_shape={m:1+2*(v-low)/(high-low) for m,v in shape.items()}
    assert all(m%100==99 for m in shape), 'Review non-.99 SKU before applying rounded USD convention'
    def level_multi(level: int) -> int:
        if level>=125:
            step=max(x for x in basis if x<=level)
            return rounded(basis[step]['vip_16_0']/original_base)
        return next(r[2] for r in cf if r[0]==max(2,level))
    tier_rows=[]
    for r in shop:
        money=r['money'];factor=new_shape[money];unit=500000*factor
        usd=(money+1)/100  # User: every .99 price is counted as the next whole USD.
        tier_rows.append([money,usd,shape[money],factor,unit,rounded(usd*unit),r['_excel_row']])
    for r in prices:
        if r['priceType'] not in (9,17):continue
        old_mult=basis[r['level']]['vip_16_0']/original_base
        scalar=500000/original_base*level_multi(r['level'])/old_mult
        if r['priceType']==9:
            scalar*=new_shape[r['money']]/shape[r['money']]
        for v in range(16):
            field=f'vip_16_{v}';value=rounded(r[field]*scalar)
            edit('PriceCheatSheet',r['_excel_row'],openpyxl.utils.get_column_letter(v+5),r[field],value,
                 '基准50万；SKU保形100%-300%；125+等级膨胀保留dev；VIP比例保留')
    rows=[]
    for l,rec in enumerate(levels,1):
        ident=cap_id[cap[l-1]];exp=new_exp[ident];oldident=old_id(l)
        prior_spin=rec['levelUpExp'] if rec['levelUpType']==1 else rec['levelUpExp']/old_exp[oldident]
        if l==5000:
            raw=None;need=rec['levelUpExp'];method='终点保留，无下一次升级'
        elif l<6:
            raw=None;need=rec['levelUpExp'] if rec['levelUpType']==1 else rounded(prior_spin*exp)
            method='CF门槛缺口；保留CR原Spin（取整误差列示）'
        elif l<50:
            raw=cf[l][3];need=rounded(raw*3/coin_unit);method='最新到达锚点差；经验按1万金币=1点同比换算'
        else:
            raw=anchor+slope*(l+1-50);need=rounded(raw*3/coin_unit);method='L30-L50线性趋势，锚定到达L50；拟合非实测'
        assert l==5000 or 0<need<2**31
        mult=level_multi(l)
        target_spin=raw/(cap[l-1]/3) if raw is not None else None
        actual_spin=None if l==5000 else need if rec['levelUpType']==1 else need/exp
        oldcf=history['cf'][l-1] if l<=300 else None
        old_raw=oldcf[1]*oldcf[2]/3 if oldcf else None
        old_rate=basis[max(x for x in basis if x<=l)]['vip_16_0']
        rows.append([l,cap[l-1],exp,need,raw,target_spin,actual_spin,mult,rec['levelUpType'] or 0,
                     old_coin[oldident],prior_spin,old_rate,old_raw,method,
                     'CF现行倍率' if 2<=l<=124 else 'L1基准' if l==1 else 'User决定：保留dev原解锁与数值，非最新CF实测',
                     '现值保留，CF金额缺口' if l<5 else '服务器区间推定' if l==5 else 'CF实测金额' if l<=50 else '尾部拟合/可用档位取下界'])
    # User: restore original difficulty after the crossing in the 270s.
    # Compare USD consumption, not raw EXP numbers across different point/currency units.
    restore_from=next(r[0] for r in rows[269:279]
                      if r[6]*r[1]/(500000*r[7]) < r[10]*r[9]/r[11])
    for row,rec in zip(rows,levels):
        if restore_from<=row[0]<5000:
            old_gross=row[10]*row[9]/row[11]
            row[3]=rounded(old_gross*(500000*row[7])/coin_unit)
            row[6]=row[3]/row[2]
            row[13]=f'L{restore_from}+恢复原dev美元消耗；换算为小整数经验，非CF拟合'
            assert 0<row[3]<2**31
        edit('LevelCfg',rec['_excel_row'],'C',rec['levelUpExp'],row[3],row[13])
    known_checks=[]
    for row in rows[5:49]:
        assert row[1]==cf[row[0]-1][1]
        assert math.isclose(row[5],row[6],rel_tol=1e-12)
        assert math.ceil(row[5])==math.ceil(row[6])
        known_checks.append(row[0])
    overview=[1]+[i+1 for i in range(1,5000) if cap[i]!=cap[i-1]]
    data={'revision':receipt['revision'],'scale':scale,'coin_unit':coin_unit,'base_coins_usd':500000,'cr_loss':.05,'cf_loss':.15,
          'anchor_exp':anchor,'exp_slope':slope,'bet_anchor':bet_anchor,'bet_slope':bet_slope,
          'rows':rows,'tiers':tier_rows,'changes':changes,'unlock_rows':output,'unlock_sheet':tables['SlotsCasinoBetUnlock']['sheet'],
          'unlock_sources':unlock_notes,'overview':overview,'known_transition_start_levels':known_checks,
          'source_exp_ratio_conflicts':live['evidence']['exp_ratio_conflicts'],
          'restore_from_level':restore_from,
          'exp_ratio_disposition':'User decision: fit to base Bet/3 and round in the common integer EXP unit; source exceptions retained, non-blocking',
          'assumptions':['User: .99 prices count as next whole USD','CR EXP=CF EXP*3/10000; min Bet gives1EXP',f'L51 to departure L{restore_from-1}: EXP anchored linear 30-50',f'L{restore_from}+ restore source dev USD cost','Bet cadence25 after50; floor to available tier',
                         'level inflation125+ preserved dev by User','CR5% / CF15% historical net loss models unchanged']}
    save(out/'alignment-inputs.json',data)
    with (out/'configuration_diff.csv').open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f);w.writerow(['file','sheet','cell','before','after','reason'])
        for name,items in changes.items():
            for item in items:w.writerow([name+'.xlsx',*item])
    return {'revision':receipt['revision'],'known_transitions_equal':len(known_checks),'level_rows':len(rows),
            'coins_per_EXP':coin_unit,'restore_from_level':restore_from,'int32':True,'candidate_cells':{k:len(v) for k,v in changes.items()},
            'ordinary_unlock_rows':sum(r[1]==0 for r in output),'svn_submitted':False}


def verify(out: Path, root: Path | None = None) -> dict:
    d=json.loads((out/'alignment-inputs.json').read_text(encoding='utf-8'))
    candidate=root or out/'candidate'
    if root:
        from svn_submit import locate_svn
        receipt=json.loads((out/'source.local.json').read_text(encoding='utf-8'))
        def svn_xml(*args: str) -> ET.Element:
            result=subprocess.run([str(locate_svn()),*args],capture_output=True)
            assert result.returncode==0, 'SVN identity/freshness read failed'
            return ET.fromstring(result.stdout)
        entry=svn_xml('info','--xml',str(root)).find('entry')
        assert entry.findtext('url')==receipt['url'] and entry.findtext('repository/uuid')==receipt['uuid']
        assert entry.findtext('relative-url')=='^/x_proj_share/dev/ExcelConfigExport/Excel'
        for item in receipt['files']:
            e=svn_xml('info','--xml',receipt['url']+'/'+item['file']+'@HEAD').find('entry')
            assert e.find('commit').get('revision')==str(item['last_changed_revision']), 'Related dependency changed; stop before commit'
        status=svn_xml('status','--xml','--ignore-externals',str(root)).findall('.//entry')
        assert {Path(e.get('path')).name for e in status}=={name+'.xlsx' for name in FILES}
        for e in status:
            s=e.find('wc-status')
            assert s.get('item')=='modified' and s.get('props') in ('normal','none') and s.get('tree-conflicted')!='true'
    counts={}
    for name in FILES:
        path=candidate/(name+'.xlsx')
        if not root:sanitize(path.parent,path.name)
        a=grid(out/'source'/(name+'.xlsx'));b=grid(path);assert a.keys()==b.keys()
        if name=='SlotsCasinoBetUnlock':
            sheet=d['unlock_sheet'];assert b[sheet][:4]==a[sheet][:4]
            assert b[sheet][4:]==d['unlock_rows']
            assert [r for r in b[sheet][4:] if r[1]==1]==[r for r in a[sheet][4:] if r[1]==1]
            old={tuple(r[:3]):r for r in a[sheet][4:]};new={tuple(r[:3]):r for r in b[sheet][4:]}
            for key in old.keys()&new.keys():assert old[key]==new[key]
            added=new.keys()-old.keys();removed=old.keys()-new.keys()
            with (out/'unlock_diff.csv').open('w',encoding='utf-8-sig',newline='') as handle:
                writer=csv.writer(handle);writer.writerow(['change','level','highroller','betId','fields'])
                writer.writerows([['add',*key,json.dumps(new[key][3:])] for key in sorted(added)])
                writer.writerows([['remove',*key,json.dumps(old[key][3:])] for key in sorted(removed)])
            counts[name]={'rows':len(d['unlock_rows']),'added_keys':len(added),'removed_keys':len(removed),
                          'retained_keys_unchanged':True,'highroller_rows_unchanged':True};continue
        edits={(s,c):new for s,c,_,new,_ in d['changes'][name]}
        n=0
        for sheet,rows in a.items():
            assert len(rows)==len(b[sheet])
            for ri,(ar,br) in enumerate(zip(rows,b[sheet]),1):
                assert len(ar)==len(br)
                for ci,(old,new) in enumerate(zip(ar,br),1):
                    cell=f'{openpyxl.utils.get_column_letter(ci)}{ri}'
                    if (sheet,cell) in edits:assert new==edits[(sheet,cell)];n+=1
                    else:assert new==old,(name,sheet,cell)
        assert n==len(edits);counts[name]={'changed_cells':n,'non_target_cells_unchanged':True}
    bet=read_source(candidate/'SlotsCasinoBetList.xlsx')['records'];byid={r['levelId']:r for r in bet}
    unlock=d['unlock_rows'];normal=[r for r in unlock if r[1]==0]
    lev=read_source(candidate/'LevelCfg.xlsx')['records']
    for row,rec in zip(d['rows'],lev):
        l=row[0];stage=max(r[0] for r in normal if r[0]<=l)
        cap=max(byid[r[2]]['bet2'] for r in normal if r[0]==stage)
        assert cap==row[1] and rec['levelUpExp']==row[3]
        if l<5000:assert 0<row[3]<2**31
    # Existing companion coin multipliers follow base Bet after native recalc.
    for r in bet:
        for field,factor in [('bet3',1.5),('bet33',1.65),('bet9',.9),('bet35',1.75)]:
            assert math.isclose(r[field],r['bet2']*factor,abs_tol=1e-5)
    price=read_source(candidate/'PriceCheatSheet.xlsx')['records']
    original_prices=read_source(out/'source/PriceCheatSheet.xlsx')['records']
    for old,new in zip(original_prices,price):
        if old['priceType'] in (9,17):
            for vi in range(16):
                # At most half a coin per independently rounded output, not a changed VIP multiplier.
                assert abs(new[f'vip_16_{vi}']*old['vip_16_0']-old[f'vip_16_{vi}']*new['vip_16_0']) <= (old['vip_16_0']+old[f'vip_16_{vi}'])/2
    pmap={(r['money'],r['priceType'],r['vipType'],r['level']):r for r in price}
    for money,usd,_,factor,rate,coins,_ in d['tiers']:
        r=pmap[(money,9,1,0)];assert r['vip_16_0']==coins
        assert math.isclose(coins/usd,rate,rel_tol=1e-10)
        assert 1<=factor<=3
    assert d['tiers'][0][4]==500000 and d['tiers'][-1][4]==1500000
    assert pmap[(99,9,1,0)]['vip_16_0']==pmap[(100,17,1,0)]['vip_16_0']==500000
    assert all(usd==(money+1)/100 for money,usd,*_ in d['tiers'])
    for r in price:
        if r['priceType']==17 and r['money']==100 and r['vipType']==1:
            l=max(1,r['level']);assert r['vip_16_0']==500000*d['rows'][l-1][7]
            old=next(x for x in original_prices if (x['money'],x['priceType'],x['vipType'],x['level'])==(100,17,1,r['level']))
            if r['level']>=125:assert r['vip_16_0']*2==old['vip_16_0']
    book=out/BOOK;sanitize(out,BOOK)
    f=openpyxl.load_workbook(book,data_only=False);v=openpyxl.load_workbook(book,data_only=True)
    assert [s.title for s in f if s.sheet_state=='visible']==['候选说明','等级_概览','等级_明细','商城档位']
    for s in f:
        assert s.freeze_panes is None
        for row in s:
            for cell in row:
                if cell.data_type=='f':assert v[s.title][cell.coordinate].value is not None and v[s.title][cell.coordinate].data_type!='e'
    for i,row in enumerate(d['rows'],8):
        s=v['等级_明细'];assert s.cell(i,1).value==row[0] and s.cell(i,2).value==row[1]
        if row[6] is not None:
            assert math.isclose(s.cell(i,3).value,row[6],rel_tol=1e-10)
            gross=row[6]*row[1]/(500000*row[7]);assert math.isclose(s.cell(i,4).value,gross,rel_tol=1e-10)
            assert math.isclose(s.cell(i,5).value,gross*.05,rel_tol=1e-10)
            assert math.isclose(s.cell(i,6).value,gross*.15,rel_tol=1e-10)
    restored=[]
    for row in d['rows'][d['restore_from_level']-1:4999]:
        old=row[10]*row[9]/row[11];actual=row[6]*row[1]/(500000*row[7])
        assert abs(actual-old)<=d['coin_unit']/(500000*row[7])/2+1e-9
        restored.append((actual-old,(actual-old)/old))
        excel_row=row[0]+7
        assert math.isclose(v['等级_明细'].cell(excel_row,14).value,old,rel_tol=1e-10)
        assert math.isclose(v['等级_明细'].cell(excel_row,15).value,actual-old,abs_tol=1e-9)
    assert len(f['等级_概览']._charts)==4 and len(f['商城档位']._charts)==1
    assert not f._external_links
    f.close();v.close()
    old300=d['rows'][298]
    tail=d['rows'][49:d['restore_from_level']-1]
    rounding={'max_abs_theoretical_spin_error':max(abs(r[5]-r[6]) for r in tail),
              'ceil_spin_difference_count':sum(math.ceil(r[5])!=math.ceil(r[6]) for r in tail),
              'max_abs_ceil_spin_difference':max(abs(math.ceil(r[5])-math.ceil(r[6])) for r in tail)}
    result={'status':'User Review / candidate only','base_revision':d['revision'],'known_transitions_equal':44,
            'int32_pass':True,'coins_per_EXP':d['coin_unit'],'5000_terminal_preserved':True,'files':counts,
            'sku_count':30,'base_rate_endpoints':[500000,1500000],'SKU_USD_rule':'.99 rounds up to whole USD',
            'shop_and_linear_base_match':True,'tail_rounding':rounding,'VIP_not_edited':True,
            'all_5000_max_bets_validated':True,'formula_errors':0,'external_links':0,'frozen_sheets':0,
            'tail_target_vs_old_arrive300_ratio':old300[4]/old300[12],
            'restored_difficulty':{'from_level':d['restore_from_level'],'level_count':len(restored),
                'max_abs_USD_rounding_error':max(abs(r[0]) for r in restored),
                'max_abs_relative_rounding_error':max(abs(r[1]) for r in restored)},
            'source_exp_ratio_conflicts':d['source_exp_ratio_conflicts'],
            'exp_ratio_disposition':d['exp_ratio_disposition'],'SVN_submitted':False}
    save(out/('precommit-validation.json' if root else 'validation.json'),result);return result


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('action',choices=['prepare','verify'])
    p.add_argument('--base',type=Path);p.add_argument('--out',required=True,type=Path);p.add_argument('--root',type=Path);a=p.parse_args()
    print(json.dumps(prepare(a.base,a.out) if a.action=='prepare' else verify(a.out,a.root),ensure_ascii=False))
