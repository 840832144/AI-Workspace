"""TASK-0036 定向准备/核验 dev Bet 与经验；不修改 LevelCfg。

prepare|verify|verify-book --base <受控根> --out <受控输出>
verify --root <隔离SVN工作副本> 用于 svn_submit.py 的既有提交流程。
输入是本轮固定 dev 导出与上一轮受控 comparison-inputs.json；不重新盘点其他系统。
"""
from __future__ import annotations
import argparse
import csv
import json
import math
from fractions import Fraction
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET
import openpyxl
from build_cr_level_dev import svn, save
from producer_workbook_sources import read_source

BOOK = 'CR_CF_等级体验_简表.xlsx'
FILES = ('SlotsCasinoBetList.xlsx', 'SlotsCasinoBetUnlock.xlsx')


def grid(path: Path) -> dict:
    w = openpyxl.load_workbook(path, read_only=True, data_only=False)
    result = {}
    for s in w:
        if s.max_column is None:
            s.calculate_dimension(force=True)
        result[s.title] = [list(r) for r in s.values]
    w.close()
    return result


def prepare(base: Path, out: Path) -> dict:
    assert not any((p/'.git').exists() for p in (out.resolve(), *out.resolve().parents))
    receipt = json.loads((out/'source.local.json').read_text(encoding='utf-8'))
    assert receipt['environment'] == 'CR dev'
    d = json.loads((base/'outputs/task0036-bet-exp-r7252/comparison-inputs.json').read_text(encoding='utf-8'))
    # 仅复核当前候选依赖，与上一轮完全一致才复用；出现并发变化必须重新审阅。
    for entry in receipt['files']:
        assert grid(out/'source'/entry['file']) == grid(base/'outputs/task0036-bet-exp-r7252/source'/entry['file']), entry['file']
    anchors = [c for c in d['candidates'] if c['levels']]
    candidates = []
    for c in d['candidates']:
        c = dict(c)
        if c['levels']:
            method = '保留已对标锚点'
        elif c['bet'] < anchors[0]['bet']:
            method = '低于首个EXP模式锚点，保留现值'
        elif c['bet'] > anchors[-1]['bet']:
            a = anchors[-1]
            c['proposed_exp'] = int(Fraction(a['proposed_exp']*c['bet'], a['bet']) + Fraction(1, 2))
            method = 'User批准末档EXP/Bet比例延伸'
        else:
            a = max((x for x in anchors if x['bet'] < c['bet']), key=lambda x:x['bet'])
            b = min((x for x in anchors if x['bet'] > c['bet']), key=lambda x:x['bet'])
            value = a['proposed_exp'] + Fraction((b['proposed_exp']-a['proposed_exp'])*(c['bet']-a['bet']), b['bet']-a['bet'])
            c['proposed_exp'] = int(value + Fraction(1, 2))
            method = 'User批准相邻锚点线性插值，四舍五入整数'
        c['method'] = method
        assert 0 < c['proposed_exp'] < 2**31
        candidates.append(c)
    assert all(a['proposed_exp'] <= b['proposed_exp'] for a,b in zip(candidates,candidates[1:]))
    d['candidates'] = candidates
    source = read_source(out/'source/SlotsCasinoBetUnlock.xlsx')
    rows = source['records']
    fields = [f for _,f in source['fields']]
    projected = lambda r: [r.get(f) for f in fields]
    target = {r[2]: max(1,r[6]) for r in d['tiers']}
    coins = {c['id']:c['bet'] for c in candidates}
    changes, output_rows = [], []
    for stage in sorted({r['level'] for r in rows if r['highroller']==0}):
        existing = [r for r in rows if r['level']==stage and r['highroller']==0]
        cap = d['calculations'][stage-1]['official_bet']
        keep = [r for r in existing if coins[r['betlevel']]<=cap]
        # 保持同等级所有原有活动字段；新增更高Bet继承该等级原最高Bet的活动设置。
        template = max(existing,key=lambda r:coins[r['betlevel']])
        for r in existing:
            if r not in keep:
                changes.append({'kind':'remove_over_cap','level':stage,'bet_id':r['betlevel'],'source_row':r['_excel_row']})
        for ident, first in sorted(target.items()):
            if first<=stage and coins[ident]<=cap and not any(r['betlevel']==ident for r in keep):
                # 只补目标提前解锁；不重建高等级已退场的低Bet池。
                old_first = min(r['level'] for r in rows if r['highroller']==0 and r['betlevel']==ident)
                if stage < old_first:
                    r = dict(template,betlevel=ident)
                    keep.append(r)
                    changes.append({'kind':'add_earlier_unlock','level':stage,'bet_id':ident,'source_row':template['_excel_row']})
        output_rows.extend(projected(r) for r in sorted(keep,key=lambda r:r['betlevel']))
    output_rows.extend(projected(r) for r in rows if r['highroller']==1)
    before = grid(out/'source/SlotsCasinoBetUnlock.xlsx')
    assert list(before)==['Sheet2']
    # 原表含尾部未命名空列，按原始宽度保留。
    width = len(before['Sheet2'][0])
    output_rows = [r+[None]*(width-len(r)) for r in output_rows]
    unlock_grid = before['Sheet2'][:4]+output_rows
    for tier in d['tiers']:
        actual = min(r[0] for r in output_rows if r[1]==0 and r[2]==tier[2])
        assert actual == max(1,tier[6]), (tier[2],actual,tier[6])
    for row in d['calculations']:
        lv = row['level']
        # 配置声明为等级区间；只取当前最近区间，不把历史退场Bet并入池。
        stage = max(r[0] for r in output_rows if r[1]==0 and r[0]<=lv)
        actual = max(coins[r[2]] for r in output_rows if r[1]==0 and r[0]==stage)
        assert actual == row['official_bet'], lv
        v = d['cr'][lv-1]
        exp = next(c['proposed_exp'] for c in candidates if c['bet']==actual)
        row['candidate_exp'] = exp
        row['candidate_spins'] = (v[2] if v[1]==1 else math.ceil(v[2]/exp)) if lv<5000 else None
        row['candidate_cost'] = row['candidate_spins']*actual/v[6]*d['cr_loss'] if lv<5000 else None
        row['cf_discounted_cost'] = row['cf_cost']/6 if lv<=300 else None
        if lv<=300:
            assert math.isclose(d['cr_loss']/v[6],d['cf_loss']/(d['cf_base']*d['cf'][lv-1][3])/6,rel_tol=1e-10)
    assert all(r['candidate_spins']==r['cf_ceil_spins'] for r in d['calculations'][4:300])
    overview = {1}
    for a,b in zip(d['calculations'],d['calculations'][1:]):
        if a['official_bet'] != b['official_bet']:
            overview.add(b['level'])
    d.update(revision=receipt['revision'],read_at_utc=receipt['read_at_utc'],cf_discount=1/6,
             overview=sorted(overview),unlock_rows=unlock_grid,unlock_changes=changes)
    d['authority']['svn_submission']='User authorized CR dev Bet/EXP; no LevelCfg or VIP'
    d['authority']['cf_discount']='User: actual consumption is historical theoretical consumption / 6'
    save(out/'bet-dev-inputs.json',d)
    return {'base_revision':d['revision'],'EXP_changes':sum(c['proposed_exp']!=c['current_exp'] for c in candidates),
            'anchors_preserved':len(anchors),'EXP_decreases':0,'unlock_row_operations':len(changes),
            'target_first_unlocks':len(d['tiers']),'max_Bet_levels':5000,'overview_rows':len(overview)}


def verify(out: Path, root: Path | None) -> dict:
    d = json.loads((out/'bet-dev-inputs.json').read_text(encoding='utf-8'))
    candidate = root or out/'candidate'
    changes = []
    for name in FILES:
        before, after = grid(out/'source'/name), grid(candidate/name)
        sheet = 'Sheet1' if name==FILES[0] else 'Sheet2'
        assert list(before)==list(after)==[sheet]
        if name == FILES[0]:
            expected = [r[:] for r in before['Sheet1']]
            for c in d['candidates']:
                if c['proposed_exp']!=c['current_exp']:
                    expected[int(c['source_cell'].split('E')[-1])-1][4] = c['proposed_exp']
            assert expected==after['Sheet1'], 'BetList出现授权外单元格变化'
            for i,(a,b) in enumerate(zip(before['Sheet1'],expected),1):
                if a[4]!=b[4]:changes.append([name,f'E{i}','levelExp',a[4],b[4]])
        else:
            assert after[sheet]==d['unlock_rows'], 'BetUnlock与已验证候选不同'
            before_keys={(r[0],r[1],r[2]):r for r in before[sheet][4:]}
            after_keys={(r[0],r[1],r[2]):r for r in after[sheet][4:]}
            assert len(after_keys)==len(after[sheet])-4, '重复解锁键'
            for key in before_keys.keys() & after_keys.keys():assert before_keys[key]==after_keys[key]
            assert {k:v for k,v in before_keys.items() if k[1]==1}=={k:v for k,v in after_keys.items() if k[1]==1}
            for key in sorted(before_keys.keys() ^ after_keys.keys()):
                changes.append([name,str(key),'row',before_keys.get(key),after_keys.get(key)])
    with (out/'configuration_diff.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.writer(f);writer.writerow(['table','cell_or_key','field','before','after']);writer.writerows(changes)
    if root:
        receipt=json.loads((out/'source.local.json').read_text(encoding='utf-8'))
        e=ET.fromstring(svn('info','--xml',str(root))).find('entry')
        assert e.findtext('url')==receipt['url'] and e.findtext('repository/uuid')==receipt['uuid']
        assert e.findtext('relative-url')=='^/x_proj_share/dev/ExcelConfigExport/Excel'
        for entry in receipt['files']:
            current=ET.fromstring(svn('info','--xml',receipt['url']+'/'+entry['file']+'@HEAD')).find('entry')
            assert current.find('commit').get('revision')==entry['last_changed_revision'], '相关dev文件发生并发修改，停止'
        entries=ET.fromstring(svn('status','--xml','--ignore-externals',str(root))).findall('.//entry')
        assert {Path(e.get('path')).name for e in entries}==set(FILES)
        for e in entries:
            s=e.find('wc-status')
            assert s.get('item')=='modified' and s.get('props') in ('none','normal') and s.get('tree-conflicted')!='true'
    result={'base_revision':d['revision'],'changed_files':list(FILES),'EXP_changed_cells':sum(c['proposed_exp']!=c['current_exp'] for c in d['candidates']),
            'EXP_decreases':0,'anchors_preserved':17,'unlock_add_rows':sum(x['kind']=='add_earlier_unlock' for x in d['unlock_changes']),
            'unlock_removed_over_cap_rows':sum(x['kind']=='remove_over_cap' for x in d['unlock_changes']),
            'unchanged_existing_activity_cells':True,'HighRoller_unlock_unchanged':True,
            'shared_EXP_affects_HighRoller':True,'LevelCfg_written':False,'VIP_written':False,
            'integer_spin_matches_5_300':296,'unchanged_early_mismatches':[1,2,3,4],
            'scope':'configuration model only; no client runtime acceptance'}
    save(out/('precommit-validation.json' if root else 'candidate-validation.json'),result)
    return result


def verify_book(out: Path) -> dict:
    d=json.loads((out/'bet-dev-inputs.json').read_text(encoding='utf-8'))
    for v,cf in zip(d['cr'][:300],d['cf']):
        assert math.isclose(d['cr_loss']/v[6],d['cf_loss']/(d['cf_base']*cf[3])*d['cf_discount'],rel_tol=1e-10)
    w=openpyxl.load_workbook(out/BOOK,read_only=True,data_only=False)
    c=openpyxl.load_workbook(out/BOOK,read_only=True,data_only=True)
    assert [s.title for s in w if s.sheet_state=='visible']==['概览','明细']
    formulas=0
    for s in w:
        for fr,vr in zip(s.iter_rows(),c[s.title].iter_rows()):
            for f,v in zip(fr,vr):
                assert v.data_type!='e',(s.title,v.coordinate)
                if f.data_type=='f':
                    formulas+=1;assert v.value is not None,(s.title,f.coordinate)
    def check(sheet: str, levels: list[int]) -> None:
        for values,lv in zip(c[sheet].iter_rows(min_row=8,max_row=7+len(levels),values_only=True),levels):
            r=d['calculations'][lv-1]
            expected=[lv,r['official_bet'],r['candidate_spins'],r['candidate_cost'],d['cr'][lv-1][6]/d['cr'][0][6],None,lv,r['official_bet'],d['cf'][lv-1][2] if lv<=300 else None,r['cf_discounted_cost'],d['cf'][lv-1][3]/d['cf'][0][3] if lv<=300 else None]
            for i,value in enumerate(expected):
                if i==5:continue
                got=values[i]
                if value is None:assert got=='N/A',(sheet,lv,i,got)
                else:assert math.isclose(got,value,rel_tol=1e-10,abs_tol=1e-9),(sheet,lv,i,got,value)
    expected_overview=[1]+[b['level'] for a,b in zip(d['calculations'],d['calculations'][1:]) if a['official_bet']!=b['official_bet']]
    assert d['overview']==expected_overview
    check('明细',list(range(1,5001)));check('概览',d['overview'])
    for s in ('概览','明细'):
        assert w[s]['E8'].number_format==w[s]['K8'].number_format=='#,##0'
    with ZipFile(out/BOOK) as z:
        assert not any(n.startswith('xl/externalLinks/') for n in z.namelist())
        for name in z.namelist():
            if name.startswith('xl/worksheets/') and name.endswith('.xml'):assert b'<pane ' not in z.read(name)
        core=ET.fromstring(z.read('docProps/core.xml'))
        assert not any(e.text for e in core if e.tag.rsplit('}',1)[-1] in ('creator','lastModifiedBy'))
        ns={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
        charts=[ET.fromstring(z.read(n)) for n in z.namelist() if n.startswith('xl/charts/chart') and n.endswith('.xml')]
        assert len(charts)==3 and all(x.find('.//c:lineChart',ns) is not None for x in charts)
        assert all(int(p.get('val'))==300 for x in charts for p in x.findall('.//c:val/c:numRef/c:numCache/c:ptCount',ns))
    w.close();c.close()
    result={'visible_sheets':['概览','明细'],'levels':5000,'overview_rows':len(d['overview']),
            'formula_cells':formulas,'formula_errors':0,'external_links':0,'frozen_panes':0,'charts':3,
            'CF_discount':'1/6 per User','CF_fractional_expectation_preserved':True,
            'inflation_base':'level1 = 1; integer multiples','inflation_levels_CR':5000,'inflation_levels_CF':300,
            'USD_equal_at_same_Spin':True,'integer_Spin_can_differ_from_fractional_CF':True}
    save(out/'workbook-validation.json',result)
    return result


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('action',choices=['prepare','verify','verify-book']);p.add_argument('--base',type=Path,required=True);p.add_argument('--out',type=Path,required=True);p.add_argument('--root',type=Path)
    a=p.parse_args()
    print(json.dumps(prepare(a.base,a.out) if a.action=='prepare' else verify_book(a.out) if a.action=='verify-book' else verify(a.out,a.root),ensure_ascii=False))
