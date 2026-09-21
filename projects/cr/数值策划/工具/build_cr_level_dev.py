"""TASK-0036: CF1–300同级目标，后续趋势拟合；只允许dev LevelCfg门槛变更。

python build_cr_level_dev.py prepare|verify --base <controlled-root> --out <controlled-output>
verify --root <isolated-svn-wc> 可作为svn_submit.py的定向validation_command。
"""
from __future__ import annotations
import argparse
import csv
import json
import math
from pathlib import Path
import subprocess
import xml.etree.ElementTree as ET
import openpyxl
from producer_workbook_sources import read_source

BOOK = 'CR_等级逐级对标及尾部拟合_dev.xlsx'


def save(path: Path, data: object) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')


def svn(*args: str) -> bytes:
    r = subprocess.run(['svn', *args], capture_output=True)
    if r.returncode:
        raise RuntimeError('SVN定向核对失败；内部响应未输出')
    return r.stdout


def prepare(base: Path, out: Path) -> dict:
    old = json.loads((base/'outputs/task0036-tuning-pop-cf-20260921/tuning-inputs.json').read_text(encoding='utf-8'))
    receipt = json.loads((out/'dev-source.local.json').read_text(encoding='utf-8'))
    assert receipt['environment'] == 'dev'
    source = out/'dev-before'
    levels = read_source(source/'LevelCfg.xlsx')['records']
    cf = old['cf']
    assert [r[0] for r in cf] == list(range(1, 301))
    assert [r['level'] for r in levels] == list(range(1, 5001))
    # 只校验本次升级计算依赖；活动显示/掉落字段不是升级成本输入，且不修改这些表。
    specs = json.loads((base/'producer-dashboard-20260920/export-receipt.json').read_text(encoding='utf-8-sig'))
    paths = {r['table']: r['file'] for r in specs['files']}
    columns = {'LevelCfg': ['level','levelUpType','levelUpExp'],
               'SlotsCasinoBetList': ['levelId','bet2','levelExp'],
               'SlotsCasinoBetUnlock': ['highroller','level','betlevel'],
               'PriceCheatSheet': ['level','vip_16_0']}
    freshness = {}
    for name, fields in columns.items():
        a = read_source(base/'producer-dashboard-20260920/source/r7013'/paths[name])['records']
        b = read_source(source/(name+'.xlsx'))['records']
        if name == 'PriceCheatSheet':
            a, b = [[r for r in rows if (r['money'],r['priceType'],r['vipType']) == (100,17,1)] for rows in (a,b)]
        project = lambda rows: [tuple(r.get(k) for k in fields) for r in rows]
        assert project(a) == project(b), f'升级依赖变化，不能直接提交: {name}'
        freshness[name] = {'fields': fields, 'rows': len(a), 'equal': True}
    anchor = cf[-1][1]*cf[-1][3]
    window = cf[249:300]
    slope = sum((r[0]-300)*(r[1]*r[3]-anchor) for r in window)/sum((r[0]-300)**2 for r in window)
    assert slope > 0
    assert max(abs(r[1]*r[3]-anchor-slope*(r[0]-300)) for r in window)/anchor < 1e-10
    rows, errors = [], []
    for r, current in zip(old['levels'], levels):
        lv, mode, exp_old, exp_spin, bet, cost_old, reward, _, _, inflation, _ = r
        target = cf[lv-1][1] if lv <= 300 else (anchor+slope*(lv-300))/inflation
        terminal = lv == 5000
        spins = max(1, math.ceil(round(target/(bet*.05), 10)))
        need = current['levelUpExp'] if terminal else spins if mode == 1 else math.ceil(round(spins*exp_spin, 10))
        if not terminal:
            assert isinstance(need, int) and 0 < need < 2**31
        realized = 'N/A' if terminal else spins*bet*.05
        delta = 'N/A' if terminal else realized/target-1
        if lv <= 300: errors.append(delta)
        rows.append([lv,mode,exp_old,exp_spin,bet,inflation,cost_old,reward,
                     'N/A' if terminal else target,'N/A' if terminal else spins,need,realized,delta,
                     '同级CF原值' if lv<=300 else '终点，无下一次升级' if terminal else '尾部拟合',current['_excel_row']])
    data = {'source_revision': receipt['revision'], 'source': str(source/'LevelCfg.xlsx'),
            'anchor': anchor, 'slope': slope, 'window': [250,300], 'cf':cf, 'rows': rows,
            'first300_actual_exact': sum(abs(x)<1e-9 for x in errors), 'max_rounding_relative': max(errors)}
    save(out/'level-inputs.json', data)
    save(out/'freshness-validation.json', {'result':'upgrade dependencies equal', 'tables':freshness,
        'preserved_differences':'dev活动type5/6解锁字段不同；BetList空drop字段结构不同；均不修改'})
    return {'prepared':5000,'identity_targets':300,'fitted_transitions':4699,'VIP':'held','int32':'passed'}


def verify(base: Path, out: Path, root: Path | None) -> dict:
    d = json.loads((out/'level-inputs.json').read_text(encoding='utf-8'))
    candidate = (root if root else out/'candidate')/'LevelCfg.xlsx'
    before = openpyxl.load_workbook(d['source'], read_only=True, data_only=False)
    after = openpyxl.load_workbook(candidate, read_only=True, data_only=False)
    assert before.sheetnames == after.sheetnames
    changes = []
    for a,b in zip(before,after):
        # Artifact导出可省略dimension；先按实际单元格恢复只读迭代边界。
        if b.max_column is None: b.calculate_dimension(force=True)
        assert a.max_column==b.max_column and a.max_row==b.max_row
        ar,br = list(a.values),list(b.values)
        assert len(ar)==len(br)
        for ri,(av,bv) in enumerate(zip(ar,br),1):
            assert len(av)==len(bv)
            for ci,(x,y) in enumerate(zip(av,bv),1):
                if x == y: continue
                assert a.title=='Sheet1' and ci==3 and 5<=ri<=5003
                expected = d['rows'][ri-5]
                assert y==expected[10] and av[0]==expected[0]
                changes.append([a.title,f'C{ri}',av[0],'levelUpExp',x,y])
    before.close();after.close()
    assert len(changes)==4999
    with (out/'LevelCfg_actual_diff.csv').open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f);w.writerow(['sheet','cell','level','field','before','after']);w.writerows(changes)
    if root:
        receipt=json.loads((out/'dev-source.local.json').read_text(encoding='utf-8'))
        entry=ET.fromstring(svn('info','--xml',str(root))).find('entry')
        assert entry.findtext('url')==receipt['url']
        assert entry.findtext('repository/uuid')==receipt['uuid']
        assert entry.findtext('relative-url')=='^/x_proj_share/dev/ExcelConfigExport/Excel'
        for r in receipt['files']:
            e=ET.fromstring(svn('info','--xml',receipt['url']+'/'+r['file']+'@HEAD')).find('entry')
            assert e.find('commit').get('revision')==r['last_changed_revision'], '远端升级依赖已变化，停止提交'
        entries=ET.fromstring(svn('status','--xml','--ignore-externals',str(root))).findall('.//entry')
        assert len(entries)==1
        item=entries[0];status=item.find('wc-status')
        assert Path(item.get('path')).resolve()==candidate.resolve() and status.get('item')=='modified'
        assert status.get('props') in ('none','normal') and status.get('tree-conflicted')!='true'
    else:
        values=openpyxl.load_workbook(out/BOOK,read_only=True,data_only=True)
        formulas=openpyxl.load_workbook(out/BOOK,read_only=True,data_only=False)
        for got,want in zip(values['等级_明细'].iter_rows(min_row=5,max_row=5004,values_only=True),d['rows']):
            for index,expected in [(0,want[0]),(4,want[8]),(5,want[9]),(6,want[10]),(7,want[11]),(8,want[12])]:
                assert math.isclose(got[index],expected,rel_tol=1e-10,abs_tol=1e-7) if isinstance(expected,(int,float)) else got[index]==expected
        for s in formulas:
            cached=values[s.title]
            for fr,vr in zip(s.iter_rows(),cached.iter_rows()):
                for f,v in zip(fr,vr):
                    assert v.data_type!='e', (s.title,v.coordinate)
                    assert f.data_type!='f' or v.value is not None, (s.title,v.coordinate)
        values.close();formulas.close()
    result={'source_revision':d['source_revision'],'changed_files':['LevelCfg.xlsx'],'changed_cells':len(changes),
            'identity_level_targets':300,'fitted_transitions':4699,'terminal_5000':'unchanged',
            'first300_actual_exact':d['first300_actual_exact'],'max_rounding_relative':d['max_rounding_relative'],
            'VIP':'not submitted','other_fields':'unchanged','status':'verified WC' if root else 'verified candidate'}
    save(out/('precommit-validation.json' if root else 'level-validation.json'),result)
    return result


def main() -> None:
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('action',choices=['prepare','verify'])
    p.add_argument('--base',type=Path,required=True);p.add_argument('--out',type=Path,required=True);p.add_argument('--root',type=Path)
    a=p.parse_args();print(json.dumps(prepare(a.base,a.out) if a.action=='prepare' else verify(a.base,a.out,a.root),ensure_ascii=False))


if __name__=='__main__':main()
