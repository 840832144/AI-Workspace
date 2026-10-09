"""Read-only verification of the delivered XLSX against its controlled model."""
import json
import math
import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET
import openpyxl


def main() -> int:
    out=Path(sys.argv[1])
    model=json.loads((out/'report-model.controlled.json').read_text(encoding='utf8'))
    layout=json.loads((out/'layout.local.json').read_text(encoding='utf8'))
    file=out/f"CR_0930_线上效果复盘_{model['cutoff']}.xlsx"
    book=openpyxl.load_workbook(file,read_only=True,data_only=True)
    formula_book=openpyxl.load_workbook(file,read_only=True,data_only=False)
    errors=[];checked=0
    def check(sheet: str, cell: str, want: float | int | str) -> None:
        nonlocal checked
        got=book[sheet][cell].value
        ok=math.isclose(got,want,rel_tol=1e-10,abs_tol=1e-10) if isinstance(want,(float,int)) and isinstance(got,(float,int)) else got==want
        if not ok:errors.append(f'{sheet}!{cell}: output differs from model')
        checked+=1
    def ratio(a: float | int | None, b: float | int | None) -> float | str:return a/b if b is not None and b>0 and a is not None else 'N/A'
    for p,l in zip(model['pages'],layout):
        for s,pos in zip(p['sections'],l['sections']):
            for row,vals in enumerate(s['rows'],pos['begin']):
                kind=s['kind']
                if kind in ('ratios','econratios','percapita'):
                    d=ratio(vals[1],vals[2]);g=ratio(vals[4],vals[5])
                    check(p['name'],f'D{row}',d);check(p['name'],f'G{row}',g)
                    if kind=='ratios':
                        numeric=isinstance(d,(int,float)) and isinstance(g,(int,float))
                        check(p['name'],f'H{row}',(g-d)*100 if numeric else 'N/A')
                        check(p['name'],f'I{row}',g/d-1 if numeric and d>0 else 'N/A')
                elif kind=='adjusted':check(p['name'],f'F{row}',(vals[4]-vals[3])*100)
                elif kind in ('depth','coverage'):check(p['name'],f'F{row}',ratio(vals[4],vals[3]))
    assert book.sheetnames==[x['name'] for x in model['pages']]
    assert len(book.sheetnames)==8
    formula_count=0;cached_errors=[];missing_cache=[]
    for ws in formula_book:
        assert ws.sheet_state=='visible'
        for row in ws:
            for c in row:
                if c.data_type=='f':
                    formula_count+=1
                    v=book[ws.title][c.coordinate]
                    if v.value is None:missing_cache.append(f'{ws.title}!{c.coordinate}')
                    if v.data_type=='e':cached_errors.append(f'{ws.title}!{c.coordinate}')
    ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main','c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
    with zipfile.ZipFile(file) as z:
        external=[n for n in z.namelist() if n.startswith('xl/externalLinks/')]
        charts=[n for n in z.namelist() if re.fullmatch(r'xl/(?:drawings/)?charts/chart\d+\.xml',n)]
        assert len(charts)==5
        assert all(ET.fromstring(z.read(n)).find('.//c:lineChart',ns) is not None for n in charts)
        assert not external
        assert all(b'TargetMode="External"' not in z.read(n) for n in z.namelist() if n.endswith('.rels'))
        assert not any(ET.fromstring(z.read(n)).find('.//s:pane',ns) is not None for n in z.namelist() if re.fullmatch(r'xl/worksheets/sheet\d+\.xml',n))
        xmltext='\n'.join(z.read(n).decode('utf8') for n in z.namelist() if n.endswith('.xml') and ('worksheets/' in n or n in ('xl/sharedStrings.xml','docProps/core.xml')))
        assert not re.search(r'\b[A-Z]:[\\/]|access_token|refresh_token|#user_id|orderid',xmltext,re.I)
    assert book['00_复盘总览']['G15'].value=='N/A'
    assert book['00_复盘总览']['E15'].value is None
    heat=next(s for s in layout[2]['sections'] if s['title']=='每日注册批次的成熟留存')
    for row in range(heat['begin'],heat['end']+1):
        dt=book.worksheets[2][f'A{row}'].value
        if dt.strftime('%Y-%m-%d')>='2026-09-30':assert book.worksheets[2][f'E{row}'].value is None
    result={'passed':not(errors or cached_errors or missing_cache),'sheets':8,'editable_line_charts':len(charts),'formula_cells':formula_count,'independent_derived_checks':checked,'mismatches':errors,'formula_errors':cached_errors,'missing_formula_caches':missing_cache,'external_links':0,'frozen_panes':0,'immature_D7_preserved':True,'identity_path_token_text_check':'passed','native_Excel_WPS_acceptance':'not performed; no UI used'}
    (out/'workbook-validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps(result,ensure_ascii=False))
    assert result['passed']
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
